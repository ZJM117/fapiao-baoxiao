# -*- coding: utf-8 -*-
"""
隔离自检：SQLite 数据层（P0-1 数据层 + P0-3 审计日志）
=====================================================
重点验收「两个人同时干活不丢数据」—— 这是从单人工具走向多人使用的第一道门。

对照改造清单的验收标准：
  · 两个用户同时改不同记录不会互相覆盖        → 第三、五节
  · 同时操作同一条记录时只有一个成功并收到提示 → 第六节（乐观锁）
  · 重启后数据完整                            → 第二节末
  · 旧 Excel 迁移前自动备份                    → 第九节

全部在 %TEMP% 副本里跑，绝不碰真实台账 / 输出 / 配置。无窗口。
跑法：  python 测试\\diag_db.py
"""
import io
import os
import sys
import tempfile
import threading
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

TMP = Path(tempfile.mkdtemp(prefix="inv_db_"))

# ⚠️ 必须在 import db / ledger 之前改路径（数据库位置跟着 LEDGER_FILE 走）
import app_paths                                                    # noqa: E402
app_paths.LEDGER_FILE = TMP / "发票台账.xlsx"
app_paths.LEDGER_BAK = TMP / "台账备份"
app_paths.OUTPUT_DIR = TMP / "输出"
app_paths.WORK_DIR = TMP / "缓存"
for _d in (app_paths.LEDGER_BAK, app_paths.OUTPUT_DIR, app_paths.WORK_DIR):
    _d.mkdir(parents=True, exist_ok=True)

import openpyxl                                                     # noqa: E402
import db                                                           # noqa: E402
import ledger                                                       # noqa: E402

FAIL = []


def chk(cond, label, extra=""):
    print(("  [OK] " if cond else "  [XX] ") + label + ("" if cond else f"   → {extra}"))
    if not cond:
        FAIL.append(label)


def section(t):
    print("\n" + "=" * 66 + f"\n{t}\n" + "=" * 66)


def rec(no="", ctype="发票", date="2025-03-05", total=100.0, detail="", src="", **kw):
    d = {"no": no, "ctype": ctype, "date": date, "total": total, "detail": detail,
         "src": src or rf"E:\桌\票\{no or ctype}-{date}.pdf", "kind": "电子发票",
         "buyer": "某某公司", "seller": "某某公司", "amount": 90.0, "tax": 10.0,
         "project": "某某项目", "project_code": "1815P825000P", "warn": []}
    d.update(kw)
    return d


print(f"临时目录：{TMP}")

# ======================================================================
section("一、建库与基本往返")
# ======================================================================
db.init()
chk(db.db_path().name == "发票台账.db", "数据库跟着台账 Excel 落在同一处", db.db_path())
chk(db.count() == 0, "新库是空的")

ledger.save([{"序号": 1, "状态": "未报销", "凭证类型": "火车票", "发票号码": "10000000000000000001",
              "开票日期": "2025-03-02", "价税合计": 75.0, "文件路径": r"E:\a.pdf", "提示": ""},
             {"序号": 2, "状态": "未报销", "凭证类型": "发票", "发票号码": "10000000000000000002",
              "开票日期": "2025-03-03", "价税合计": 1100.0, "文件路径": r"E:\b.pdf", "提示": ""}],
            do_backup=False)
rs = ledger.load()
chk(len(rs) == 2, "save → load 往返 2 行", len(rs))
chk([r["序号"] for r in rs] == [1, 2], "序号是算出来的、连续", [r["序号"] for r in rs])
chk(rs[0]["价税合计"] == 75.0 and isinstance(rs[0]["价税合计"], float), "金额读回来还是数字")
chk(set(rs[0].keys()) >= set(ledger.FIELDS), "行 dict 带齐 22 个中文列名")

wb = openpyxl.load_workbook(str(app_paths.LEDGER_FILE))
chk(wb[ledger.SHEET]["A1"].value == "序号", "Excel 快照表头是「序号」")
chk(wb[ledger.SHEET].max_row == 3, "Excel 快照 3 行（表头 + 2）", wb[ledger.SHEET].max_row)

# ======================================================================
section("二、查重三条路（号码 / 指纹 / 路径）")
# ======================================================================
r = ledger.add_records([rec(no="10000000000000000001")])
chk(len(r["added"]) == 0 and len(r["in_ledger"]) == 1, "同号码再入 → 认成「台账里已有」")
chk(ledger.count() == 2, "没有重复登记", ledger.count())

ledger.update_rows([1], 状态="已报销")
r = ledger.add_records([rec(no="10000000000000000001")])
chk(len(r["dup"]) == 1, "已报销的票再入 → 标成重复报销风险")
chk("已报销" in r["dup"][0]["tip"], "提示里说清了为什么", r["dup"][0]["tip"][:30])

# 无号码的行程单：同一份内容指纹，入两次只留一条
cab = rec(no="", ctype="打车行程单", date="2025-03-07", total=13.58,
          detail="悦行优选 合肥南站 → 财智中心", src=r"E:\桌\票\cx1.pdf")
n0 = ledger.count()
r1 = ledger.add_records([cab])
r2 = ledger.add_records([dict(cab, src=r"E:\桌\票\换个文件名.pdf")])
chk(len(r1["added"]) == 1, "无号码行程单第一次入库")
chk(len(r2["added"]) == 0 and ledger.count() == n0 + 1,
    "同一份行程单（指纹相同）第二次不重复入库", f"count={ledger.count()}")

# 兜底那一路：老台账里的记录没有号码、指纹也算不出来，就靠「同一个文件路径」认出来
ledger.save([{"序号": 1, "状态": "未报销", "凭证类型": "图片",
              "文件路径": r"E:\桌\票\old.jpg", "提示": ""}], do_backup=False)
chk(ledger.count() == 1, "先铺一条只有路径的老记录")
r3 = ledger.add_records([rec(no="", ctype="图片2", date="", total=0.0, detail="",
                             src=r"E:\桌\票\old.jpg")])
chk(len(r3["added"]) == 0 and ledger.count() == 1,
    "号码和指纹都对不上时，靠同一个文件路径认成已有", r3["in_ledger"])
r4 = ledger.add_records([rec(no="", ctype="图片2", date="", total=0.0, detail="",
                             src=r"E:\桌\票\new.jpg")])
chk(len(r4["added"]) == 1, "换个路径就是新的一条")

before = ledger.count()
chk(db.init() is None and ledger.count() == before, "重启（重新建连）后数据完整", ledger.count())

# ======================================================================
section("三、并发入库：10 个人同时入库，一条都不能丢 ★")
# ======================================================================
ledger.clear_ledger()
chk(ledger.count() == 0, "先清空")

errors = []
N = 10


def worker(i):
    try:
        ledger.add_records([rec(no=f"200000000000000000{i:02d}",
                                date=f"2025-04-{i + 1:02d}", total=10.0 * (i + 1))])
    except Exception as e:                                     # noqa: BLE001
        errors.append(f"{type(e).__name__}: {e}")


ts = [threading.Thread(target=worker, args=(i,)) for i in range(N)]
for t in ts:
    t.start()
for t in ts:
    t.join()
chk(not errors, f"{N} 个线程都没报错", errors[:3])
chk(ledger.count() == N, f"台账里正好 {N} 行（旧实现这里会少几条）", ledger.count())
nos = {str(r["发票号码"]) for r in ledger.load()}
chk(len(nos) == N, "10 张票一张不缺、不重", len(nos))

wb = openpyxl.load_workbook(str(app_paths.LEDGER_FILE))
chk(wb[ledger.SHEET].max_row == N + 1,
    "Excel 快照也跟着是全部行（不是中途某个时刻的快照）", wb[ledger.SHEET].max_row)

# ======================================================================
section("四、并发入同一张票：只入一次、不报错")
# ======================================================================
ledger.clear_ledger()
SAME = "30000000000000000099"
errs2 = []


def same_worker():
    try:
        ledger.add_records([rec(no=SAME)])
    except Exception as e:                                     # noqa: BLE001
        errs2.append(f"{type(e).__name__}: {e}")


ts = [threading.Thread(target=same_worker) for _ in range(8)]
for t in ts:
    t.start()
for t in ts:
    t.join()
chk(not errs2, "8 个线程抢着入同一张票，没有异常冒出来", errs2[:3])
chk(ledger.count() == 1, "同一张票只入了一行", ledger.count())

# ======================================================================
section("五、并发改不同行：各改各的，互不覆盖")
# ======================================================================
ledger.clear_ledger()
for i in range(8):
    ledger.add_records([rec(no=f"400000000000000000{i:02d}", date=f"2025-05-{i + 1:02d}")])
chk(ledger.count() == 8, "先铺 8 行")


def upd(i):
    try:
        ledger.update_rows([i + 1], 费用类别=f"类别{i}")
    except Exception as e:                                     # noqa: BLE001
        errors.append(f"{type(e).__name__}: {e}")


ts = [threading.Thread(target=upd, args=(i,)) for i in range(8)]
for t in ts:
    t.start()
for t in ts:
    t.join()
rows = ledger.load()
cats = [r["费用类别"] for r in rows]
chk(cats == [f"类别{i}" for i in range(8)],
    "8 行各是各的类别，没有谁把谁的改动整片盖掉", cats)

# ======================================================================
section("六、同一行并发编辑：拿着过期版本号改不动")
# ======================================================================
ledger.clear_ledger()
ledger.add_records([rec(no="50000000000000000001")])
snap = db.all_rows_with_id()
rid = snap[0]["_id"]
with db.write_tx() as c:
    n1 = db.update_row_by_id(c, rid, {"费用类别": "甲改的"}, expect_rev=0)
with db.write_tx() as c:
    n2 = db.update_row_by_id(c, rid, {"费用类别": "乙改的"}, expect_rev=0)
chk(n1 == 1, "先手（版本号对得上）改成功")
chk(n2 == 0, "后手拿着同一个旧版本号 → 改不动，界面据此提示「这行已被别人改过」")
chk(ledger.load()[0]["费用类别"] == "甲改的", "最终值就是成功那一次的值",
    ledger.load()[0]["费用类别"])
with db.write_tx() as c:
    n3 = db.update_row_by_id(c, rid, {"费用类别": "甲再改"}, expect_rev=1)
chk(n3 == 1, "拿到新版本号就能继续改")

# ======================================================================
section("七、软删除 + 审计留痕")
# ======================================================================
db.set_actor("张三", "192.168.1.20")
ledger.clear_ledger()
for i in range(3):
    ledger.add_records([rec(no=f"600000000000000000{i:02d}", date=f"2025-06-{i + 1:02d}")])
n = ledger.delete_rows([2])
chk(n == 1, "删掉 1 行", n)
chk(len(ledger.load()) == 2, "界面上只剩 2 行")
with db.read() as c:
    alive = c.execute("SELECT COUNT(*) n FROM invoices WHERE is_deleted = 0").fetchone()["n"]
    dead = c.execute("SELECT COUNT(*) n FROM invoices WHERE is_deleted = 1").fetchone()["n"]
chk(alive == 2 and dead >= 1, "被删的行还在库里（软删除，审计追得回来）", f"活{alive}/删{dead}")

aud = db.list_audit(limit=20, action="删除台账行")
chk(len(aud) >= 1, "删除动作写进了审计")
chk(aud and aud[0]["username"] == "张三", "审计记下了是谁干的", aud[0]["username"] if aud else "")
chk(aud and aud[0]["ip"] == "192.168.1.20", "审计记下了从哪来", aud[0]["ip"] if aud else "")
chk(aud and "60000000000000000001" in (aud[0]["before_json"] or ""),
    "审计留了改前内容")

n = ledger.clear_ledger()
chk(n == 2, "清空台账清掉 2 行", n)
chk(db.count() == 0, "清空之后库里活行是 0")
chk(db.stats()["deleted"] >= 3, "删掉的行都留在库里", db.stats())

# ======================================================================
section("八、Excel 快照：跟库一致 + 手改不被静默覆盖")
# ======================================================================
ledger.clear_ledger()
for i in range(3):
    ledger.add_records([rec(no=f"700000000000000000{i:02d}", date=f"2025-07-{i + 1:02d}")])
chk(openpyxl.load_workbook(str(app_paths.LEDGER_FILE))[ledger.SHEET].max_row == 4,
    "快照与库一致（表头 + 3 行）")

# 模拟用户自己用 Excel 改了一个格子
wb = openpyxl.load_workbook(str(app_paths.LEDGER_FILE))
ws = wb[ledger.SHEET]
ws["P2"] = "我手工写的备注"
wb.save(str(app_paths.LEDGER_FILE))
n_before = len(list(app_paths.LEDGER_BAK.glob("*手改*")))
info = ledger.export_excel()
n_after = len(list(app_paths.LEDGER_BAK.glob("*手改*")))
chk(n_after == n_before + 1, "重导之前先把「手改过的那份」另存了备份", info)
chk(bool(info.get("backup")), "并且把备份路径回传出来了", info.get("backup"))

n2 = len(list(app_paths.LEDGER_BAK.glob("*手改*")))
ledger.export_excel()
chk(len(list(app_paths.LEDGER_BAK.glob("*手改*"))) == n2,
    "没再手改的话，不会每次都堆一份备份")

# ======================================================================
section("九、从老 Excel 迁移（迁移前自动备份）")
# ======================================================================
T2 = TMP / "迁移"
T2.mkdir(exist_ok=True)
app_paths.LEDGER_FILE = T2 / "发票台账.xlsx"
app_paths.LEDGER_BAK = T2 / "台账备份"
app_paths.LEDGER_BAK.mkdir(parents=True, exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "发票台账"
ws.append([c[0] for c in ledger.COLUMNS])
ws.append([1, "未报销", "火车票", "", "张三", "80000000000000000001", "2025-08-01", "铁路客票",
           "G1 武汉→宜昌", "某某公司", "中国铁路武汉局", 71.56, 3.44, 75.0,
           "某项目 1815P825000P", "1815P825000P", "交通费", "张三", "8月报销单",
           "2025-08-01 10:00", r"E:\a.pdf", ""])
ws.append([2, "已报销", "发票", "", "李四", "80000000000000000002", "2025-08-02", "电子发票",
           "", "某某公司", "某某酒店", 1037.74, 62.26, 1100.0,
           "某项目 1815P825000P", "1815P825000P", "住宿费", "李四", "8月报销单",
           "2025-08-02 10:00", r"E:\b.pdf", ""])
wb.save(str(app_paths.LEDGER_FILE))

chk(db.count() == 0, "迁移前这个库是空的")
res = db.migrate_from_excel(app_paths.LEDGER_FILE)
chk(res["ok"] and res["migrated"] == 2, "迁进来 2 行", res)
chk(Path(res["backup"]).exists(), "迁移前的原文件已经备份", res["backup"])
chk(app_paths.LEDGER_FILE.exists(), "原 Excel 没有被删（只加不删）")
rows = ledger.load()
chk(len(rows) == 2 and rows[1]["状态"] == "已报销", "迁进来的内容和 Excel 一致")
chk(rows[0]["价税合计"] == 75.0, "金额正确", rows[0]["价税合计"])

res2 = db.migrate_from_excel(app_paths.LEDGER_FILE)
chk(res2.get("skipped"), "再迁一次会跳过，不会把数据翻倍", res2.get("skipped"))

# 历史台账里真有「同一张票入过两次」的留痕（两个批次各入一次）
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "发票台账"
ws.append([c[0] for c in ledger.COLUMNS])


def _dup_row(seq, status, when, src):
    return [seq, status, "发票", "", "", "90000000000000000001", "2025-09-01", "电子发票", "",
            "某某公司", "甲公司", 90.0, 10.0, 100.0, "", "", "交通费", "", "", when, src, ""]


ws.append(_dup_row(1, "未报销", "2025-09-01 10:00", r"E:\d1.pdf"))
ws.append(_dup_row(2, "已报销", "2025-09-02 10:00", r"E:\d2.pdf"))
wb.save(str(app_paths.LEDGER_FILE))
res3 = db.migrate_from_excel(app_paths.LEDGER_FILE, force=True)
chk(res3["ok"] and res3["migrated"] == 2,
    "同一个发票号两行的历史台账，整批迁得进来（不因唯一索引整批失败）", res3)
chk(ledger.count() == 2, "两行都保住了，一条没丢", ledger.count())
with db.read() as c:
    n_dupe = c.execute("SELECT COUNT(*) n FROM invoices WHERE dupe = 1").fetchone()["n"]
chk(n_dupe == 1, "重复的那条被标成 dupe（保留数据、不参与唯一性约束）", n_dupe)

# ======================================================================
section("十、后台任务表")
# ======================================================================
jid = db.job_create("导入", "入库 30 个文件", total=30)
chk(bool(jid), "建任务拿到任务号")
db.job_update(jid, state=db.JOB_RUNNING, done=12)
j = db.job_get(jid)
chk(j["state"] == db.JOB_RUNNING and j["done"] == 12, "任务进度写得进去", j)
db.job_update(jid, state=db.JOB_DONE, done=30, ended_at="2025-09-16 10:00:00")
chk(db.job_get(jid)["state"] == db.JOB_DONE, "任务能收尾")
jid2 = db.job_create("导出", "生成报销单", total=1)
db.job_update(jid2, state=db.JOB_RUNNING)
n = db.job_orphans_fix()
chk(n >= 0, f"重启后能把挂着的任务收掉（本次收了 {n} 个）")
chk(db.job_get(jid2)["state"] == db.JOB_FAILED, "上一世没跑完的任务被标成失败，不会一直转圈")

# ======================================================================
section("十一、口令哈希")
# ======================================================================
h = db.hash_password("S3cret-口令")
chk(h and not h.startswith("S3cret"), "口令不是明文存的", h[:24] + "…")
chk(db.verify_password("S3cret-口令", h), "对的口令能验过")
chk(not db.verify_password("错口令", h), "错的口令验不过")
chk(db.hash_password("同一个口令") != db.hash_password("同一个口令"),
    "同一个口令每次哈希都不一样（加了盐）")
chk(db.verify_password("明文口令", "明文口令"), "老配置里的明文字令仍能登录（平滑过渡）")

db.user_create("zhangsan", "pw123456", role="biz", display_name="张三")
u = db.user_get("zhangsan")
chk(u["role"] == "biz" and u["disabled"] == 0, "建用户成功")
chk(db.can("zhangsan", "biz", "start_import"), "业务人员可以入库")
chk(not db.can("zhangsan", "biz", "clear_ledger"), "业务人员不能清空台账")
chk(db.can("admin", "admin", "clear_ledger"), "管理员什么都能干")
chk(not db.can("v", "viewer", "make_report"), "只读用户不能出单据")

# ======================================================================
print("\n" + "=" * 66)
print(f"结果：{len(FAIL) == 0 and '全部通过' or '有失败'}"
      f"（失败 {len(FAIL)} 项）")
for f in FAIL:
    print("  ✗", f)
print(f"临时目录：{TMP}")
sys.exit(1 if FAIL else 0)
