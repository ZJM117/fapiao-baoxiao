# -*- coding: utf-8 -*-
"""定向探测：高铁票 PDF / 航空行程单 / OFD 内部结构 / 用户已有 Excel 表头。只读。"""
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

B = Path(r"E:\桌\湖北报销发票整理")

PDFS = [
    B / r"报销\李强（某某公司）\济南东-烟台 李强 10.pdf",
    B / r"报销\李强（某某公司）\武汉-烟台  李强电子行程单748元.pdf",
    B / r"某某公司\待报销的发票\网上购票系统-电子发票通知\恩施\汉口至恩施174元 7月28日.pdf",
    B / r"某某公司\26429121050003258897.pdf",
    B / r"某某公司\待报销的发票\高德打车电子发票\【T3出行-23.58元-1个行程】高德打车电子行程单.pdf",
    B / r"某某公司\滴滴出行电子发票及行程报销单\滴滴出行行程报销单A.pdf",
    B / r"某某公司\待报销的发票\随州住宿费600元.pdf",
]

OFDS = [
    B / r"某某公司\待报销的发票\网上购票系统-电子发票通知\合肥\26349119423005870096.ofd",
    B / r"某某公司\待报销的发票\高德打车电子发票\【T3出行-23.58元-1个行程】高德打车电子发票.ofd",
]

XLSX = [
    B / r"报销\出差台账.xlsx",
    B / r"某某公司\报销明细-某某公司.xlsx",
    B / r"某某公司\待报销的发票\统计表.xlsx",
]


def show_pdf(p: Path):
    print("=" * 78)
    print(f"【PDF】{p.name}  {p.stat().st_size//1024} KB")
    if not p.exists():
        print("   !! 不存在")
        return
    import invparse
    # 走解析器实际用的链路（pypdf → 失败则 pdfminer）
    t1 = ""
    try:
        from pypdf import PdfReader
        t1 = "\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)
    except Exception as e:
        print(f"   pypdf 失败：{type(e).__name__}: {str(e)[:70]}")
    t2 = ""
    try:
        from pdfminer.high_level import extract_text
        t2 = extract_text(str(p)) or ""
    except Exception as e:
        print(f"   pdfminer 失败：{type(e).__name__}: {str(e)[:70]}")
    t = invparse._pdf_text(p)
    print(f"   pypdf {len(t1)} 字 / pdfminer {len(t2)} 字 / 实际采用 {len(t)} 字"
          + ("   ⚠ 图片型（无文字层）" if not t.strip() else ""))
    for ln in [x.rstrip() for x in t.splitlines()]:
        if ln.strip():
            print("   | " + ln[:150])


def show_ofd(p: Path):
    print("=" * 78)
    print(f"【OFD】{p.name}")
    if not p.exists():
        print("   !! 不存在")
        return
    with zipfile.ZipFile(p) as z:
        for n in z.namelist()[:30]:
            print(f"   - {n}  ({z.getinfo(n).file_size} B)")
        print("   --- 含 SellerName/EIid 的条目 ---")
        hit = 0
        for n in z.namelist():
            if not n.lower().endswith(".xml"):
                continue
            d = z.read(n)
            if b"SellerName" in d or b"EIid" in d:
                hit += 1
                print(f"   ★ {n}  {len(d)} B  head={d[:160]!r}")
        if not hit:
            print("   （没有）")


def show_xlsx(p: Path):
    print("=" * 78)
    print(f"【XLSX】{p.name}")
    if not p.exists():
        print("   !! 不存在")
        return
    import openpyxl
    wb = openpyxl.load_workbook(str(p), data_only=True)
    for ws in wb.worksheets:
        print(f"   --- sheet「{ws.title}」{ws.max_row} 行 × {ws.max_column} 列 ---")
        for i, row in enumerate(ws.iter_rows(max_row=8, values_only=True), 1):
            cells = [("" if c is None else str(c))[:22] for c in row]
            print(f"     {i}: " + " | ".join(cells))


for p in PDFS:
    show_pdf(p)
for p in OFDS:
    show_ofd(p)
for p in XLSX:
    show_xlsx(p)
