# -*- coding: utf-8 -*-
"""
隔离自检：台账删除/清空 + 多条件查询 + 报销单 PDF/Excel 双格式 + 生成文件删除
================================================================================
全部在 %TEMP% 副本里跑，绝不碰真实台账 / 输出 / 配置。无窗口。
跑法：  python 测试\\diag_crud.py
"""
import io
import sys
import tempfile
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

TMP = Path(tempfile.mkdtemp(prefix="inv_crud_"))

# ⚠️ 必须在 import ledger / app_web 之前改路径（它们用 from app_paths import XXX 取值）
import app_paths                                                    # noqa: E402
app_paths.LEDGER_FILE = TMP / "发票台账.xlsx"
app_paths.LEDGER_BAK = TMP / "台账备份"
app_paths.OUTPUT_DIR = TMP / "输出"
app_paths.WORK_DIR = TMP / "缓存"
app_paths.CONFIG_FILE = TMP / "config.json"
app_paths.LAUNCH_LOG = TMP / "launch.log"
app_paths.UI_PORT_FILE = TMP / "ui_port.txt"
for _d in (app_paths.LEDGER_BAK, app_paths.OUTPUT_DIR, app_paths.WORK_DIR):
    _d.mkdir(parents=True, exist_ok=True)

import app_web                                                      # noqa: E402
import ledger                                                       # noqa: E402
import report_pdf as rp                                             # noqa: E402

FAIL = []


def chk(cond, label, extra=""):
    print(("  [OK] " if cond else "  [XX] ") + label + ("" if cond else f"   → {extra}"))
    if not cond:
        FAIL.append(label)


def section(t):
    print("\n" + "=" * 66 + f"\n{t}\n" + "=" * 66)


def row(seq, status, ctype, no, date, kind, trip, seller, amt, tax, total,
        proj, cat, person, batch):
    return {"序号": seq, "状态": status, "凭证类型": ctype, "发票号码": no,
            "开票日期": date, "票种": kind, "行程/明细": trip, "购方名称": "某某公司",
            "销方名称": seller, "不含税金额": amt, "税额": tax, "价税合计": total,
            "项目/事由": proj, "项目编号": "1815P825000P", "费用类别": cat,
            "报销人": person, "报销批次": batch, "入库时间": "2026-09-15 13:00",
            "文件路径": rf"E:\桌\湖北报销发票整理\{no or 'x'}.pdf", "提示": ""}


ROWS = [
    row(1, "未报销", "火车票", "", "2025-03-02", "电子发票（铁路电子客票）",
        "G1234 武汉 → 宜昌东 2025-03-02 08:15 二等座 07车12F", "中国铁路武汉局",
        71.56, 3.44, 75.00, "恩施建始龙坪网格10千伏线路 1815P825000P", "交通费", "张三", "9月现场审计"),
    row(2, "未报销", "火车票", "", "2025-03-04", "电子发票（铁路电子客票）",
        "D2233 宜昌东 → 武汉 2025-03-04 17:40 二等座 03车05A", "中国铁路武汉局",
        61.47, 2.53, 64.00, "恩施建始龙坪网格10千伏线路 1815P825000P", "交通费", "张三", "9月现场审计"),
    row(3, "未报销", "发票", "26349119423005870096", "2025-03-05", "电子发票（普通发票）",
        "", "武汉某某酒店管理有限公司", 1037.74, 62.26, 1100.00,
        "恩施建始龙坪网格10千伏线路 1815P825000P", "住宿费", "张三", "9月现场审计"),
    row(4, "已报销", "发票", "26349119423005870097", "2025-02-10", "电子发票（专用发票）",
        "", "武汉某办公用品有限公司", 353.98, 46.02, 400.00,
        "黄石阳新富池供电所台区改造 1815P825001P", "办公费", "李四", "8月报销单"),
    row(5, "未报销", "打车行程单", "", "", "",
        "悦行优选 合肥南站 → 财智中心 13.58元", "南京领行科技股份有限公司",
        13.58, 0, 13.58, "", "交通费", "张三", "9月现场审计"),
]
ledger.save(ROWS, do_backup=False)
api = app_web.Api()

# ---------------------------------------------------------------- 查询
section("一、多条件查询（这次重点修好了「凭证类型」筛选）")
f = app_web.filter_rows(ROWS, {"ctype": "火车票"})
chk([r["序号"] for r in f] == [1, 2], "按凭证类型筛：火车票 = 2 条", f)
f = app_web.filter_rows(ROWS, {"batch": "8月报销单"})
chk([r["序号"] for r in f] == [4], "按报销批次筛：8月报销单 = 1 条", f)
f = app_web.filter_rows(ROWS, {"person": "张三"})
chk(len(f) == 4, "按报销人筛：张三 = 4 条", f)
f = app_web.filter_rows(ROWS, {"min": "100", "max": "1200"})
chk([r["序号"] for r in f] == [3, 4], "按金额区间 100~1200 = 2 条", f)
f = app_web.filter_rows(ROWS, {"kw": "G1234"})
chk([r["序号"] for r in f] == [1], "关键词搜行程明细（G1234）= 1 条", f)
f = app_web.filter_rows(ROWS, {"kw": "宜昌东"})
chk([r["序号"] for r in f] == [1, 2], "关键词搜行程（宜昌东）= 2 条（来回都命中）", f)
f = app_web.filter_rows(ROWS, {"status": "未报销", "category": "交通费"})
chk(len(f) == 3, "组合条件（未报销 + 交通费）= 3 条", f)
f = app_web.filter_rows(ROWS, {"ctype": "发票", "status": "已报销"})
chk([r["序号"] for r in f] == [4], "组合（发票 + 已报销）= 1 条", f)

led = api.get_ledger({"ctype": "火车票"})
chk(led["total"] == 2 and led["all_count"] == 5, "get_ledger 命中 2 / 全部 5", led["total"])
chk("batches" in led and "persons" in led and "ctypes" in led,
    "get_ledger 带回 batches / persons / ctypes 下拉", list(led.keys()))

# ---------------------------------------------------------------- 报销单
section("二、报销单：PDF + Excel（用户说「太简单、金额不全面」）")
meta = {"kind": "差旅费报销单", "person": "张三", "dept": "审计部",
        "reason": "项目现场审计差旅费", "note": "附原始凭证若干",
        "date": "2026年09月15日", "start": "2025-03-01", "end": "2025-03-05",
        "place": "湖北宜昌"}
seqs = [1, 2, 3, 5]
d = rp.report_data([r for r in ROWS if r["序号"] in seqs], meta)
chk(abs(d["total"] - 1252.58) < 0.01, f"合计 = ¥1,252.58（实际 {d['total']}）")
chk(abs(d["amount"] + d["tax"] - d["total"]) < 0.01,
    f"不含税({d['amount']}) + 税额({d['tax']}) = 合计({d['total']})，口径对得上")
chk(d["upper"] == "壹仟贰佰伍拾贰圆伍角捌分", f"大写金额（实际 {d['upper']}）")
chk(d["days"] == "5 天", f"出差天数自动算（实际 {d['days']}）")
chk(set(d["by_cat"].keys()) == {"交通费", "住宿费"},
    f"按费用类别分组 = 交通费/住宿费（实际 {list(d['by_cat'].keys())}）")
chk(d["by_cat"]["住宿费"]["total"] == 1100.00, "住宿费小计 = 1100.00")
chk("火车票 2 张" in d["attach"] and "合计 4 张" in d["attach"],
    f"附件清单（实际 {d['attach']}）")

res = api.make_report(meta, seqs, "both")
if res.get("error"):
    chk(False, "make_report(fmt='both') 生成成功", res["error"])
else:
    chk(len(res["outs"]) == 2, "一次出两个文件", res["outs"])
    xlsx = Path(res["outs"][1]) if res["outs"][1].endswith(".xlsx") else Path(res["outs"][0])
    pdf = Path(res["outs"][0]) if res["outs"][0].endswith(".pdf") else Path(res["outs"][1])
    chk(pdf.exists() and pdf.stat().st_size > 3000, f"PDF 已生成（{pdf.stat().st_size} 字节）")
    chk(xlsx.exists() and xlsx.stat().st_size > 4000, f"Excel 已生成（{xlsx.stat().st_size} 字节）")

    # PDF 里到底写了什么
    # ⚠️ 标题用了 letter-spacing，pypdf 抽出来的字之间会带空格，比对前把空白全部去掉
    from pypdf import PdfReader
    txt = "".join((p.extract_text() or "") for p in PdfReader(str(pdf)).pages)
    txt_flat = "".join(txt.split())
    for want in ("差旅费报销单", "费用明细", "金额汇总", "价税合计", "不含税金额", "税额",
                 "报销金额（大写）", "壹仟贰佰伍拾贰圆伍角捌分", "附件张数", "领款人",
                 "出差起止", "5天"):
        chk(want.replace(" ", "") in txt_flat, f"PDF 里有「{want}」")
    chk("G1234" in txt, "PDF 里带出了行程明细（G1234）")
    chk("1,252.58" in txt, "PDF 里金额带千分位（1,252.58）")

    # Excel 里到底写了什么
    import openpyxl
    wb = openpyxl.load_workbook(str(xlsx))
    ws = wb[wb.sheetnames[0]]
    cells = [str(c.value) for r in ws.iter_rows() for c in r if c.value is not None]
    flat = " | ".join(cells)
    for want in ("差旅费报销单", "费用明细（共 4 张凭证）", "金额汇总",
                 "报销金额（大写）", "壹仟贰佰伍拾贰圆伍角捌分", "附件清单", "财务审核",
                 "领款人"):
        chk(want in flat, f"Excel 里有「{want}」")
    # 金额是数值单元格（不是文本），千分位是 number_format 决定的显示效果
    total_cells = [c for r in ws.iter_rows() for c in r
                   if isinstance(c.value, (int, float)) and abs(float(c.value) - 1252.58) < 0.005]
    chk(total_cells and all("#,##0.00" in c.number_format for c in total_cells),
        f"Excel 里合计是数值 + 千分位格式（{total_cells[0].number_format if total_cells else '没找到'}）")
    cat_cell = [c for r in ws.iter_rows() for c in r if c.value == "住宿费"]
    chk(len(cat_cell) == 1, "Excel 金额汇总里有「住宿费」一行")
    chk("$A$1" in (ws.print_area or ""), f"Excel 设了打印区域（{ws.print_area}）")
    chk(ws.page_setup.fitToWidth == 1, "Excel 按宽度缩放到一页")

# 只要 PDF
res2 = api.make_report(dict(meta, kind="费用报销单"), seqs, "pdf")
chk(not res2.get("error") and res2["outs"][0].endswith(".pdf"),
    "fmt='pdf' 只出 PDF", res2.get("error") or res2["outs"])
res3 = api.make_report(dict(meta, kind="费用报销单"), seqs, "xlsx")
chk(not res3.get("error") and res3["outs"][0].endswith(".xlsx"),
    "fmt='xlsx' 只出 Excel", res3.get("error") or res3["outs"])

# ---------------------------------------------------------------- 生成文件历史 / 删除
section("三、生成的文件：历史记录 + 删除（送回收站）")
lo = api.list_outputs()
chk(len(lo["items"]) == 4, f"历史里记了 4 个文件（实际 {len(lo['items'])}）", lo["items"])
chk(lo["items"][0]["time"] >= lo["items"][-1]["time"], "最新的排在最前面")
chk(all("exists" in it for it in lo["items"]), "每项都带 exists 标记")
target = lo["items"][0]
before = Path(target["path"])
chk(before.exists(), "目标文件存在（准备删）")
r = api.delete_output(target["path"])
chk(r.get("ok") and not before.exists(), f"文件已删除（{r}）")
lo2 = api.list_outputs()
chk(len(lo2["items"]) == 3, f"历史里剩 3 个（实际 {len(lo2['items'])}）")
chk(all(it["path"] != target["path"] for it in lo2["items"]), "被删的那条不在历史里了")

r = api.delete_output(r"C:\Windows\System32\drivers\etc\hosts")
chk(r.get("error"), "越界的路径拒绝删除（只允许输出目录）", r)

# ---------------------------------------------------------------- 删除 / 清空台账
section("四、台账删除 / 清空")
r = api.delete_rows([2, 3])
chk(r.get("deleted") == 2, f"删掉 2 条（{r}）")
left = ledger.load()
chk(len(left) == 3, f"台账剩 3 行（实际 {len(left)}）")
chk([x["序号"] for x in left] == [1, 2, 3], "序号已重排连续", [x["序号"] for x in left])
chk(all(str(x["发票号码"]) != "26349119423005870096" for x in left), "被删的发票号不在了")
baks = list(app_paths.LEDGER_BAK.glob("发票台账_*.xlsx"))
chk(len(baks) >= 1, f"删除前自动备份了 {len(baks)} 份台账")

r = api.clear_ledger()
chk(r.get("cleared") == 3, f"清空 3 条（{r}）")
chk(ledger.load() == [], "台账已空")
import openpyxl as _ox
wb = _ox.load_workbook(str(app_paths.LEDGER_FILE))
chk(wb[ledger.SHEET]["A1"].value == "序号" and wb[ledger.SHEET].max_row == 1,
    "表头保留、没有数据行")
chk(len(list(app_paths.LEDGER_BAK.glob("发票台账_*.xlsx"))) >= 2, "清空前也备份了")

# 空台账上再删 / 再清空不能炸
chk(api.delete_rows([1]).get("deleted") == 0, "空台账删行返回 0")
chk(api.clear_ledger().get("cleared") == 0, "空台账清空返回 0")

# ---------------------------------------------------------------- 入库目录
section("六、入库目录必须跟着界面走（用户报的 bug：改了路径，扫的还是老目录）")
import json                                                         # noqa: E402

INV_XML = """<?xml version="1.0" encoding="UTF-8"?>
<EInvoice>
  <EIid>26349119423005870096</EIid>
  <IssueTime>2025-03-05</IssueTime>
  <EInvoiceType><LabelName>电子发票</LabelName></EInvoiceType>
  <GeneralOrSpecialVAT><LabelName>普通发票</LabelName></GeneralOrSpecialVAT>
  <InIssuType><LabelCode>Y</LabelCode></InIssuType>
  <TotalTax-includedAmount>1100.00</TotalTax-includedAmount>
  <TotalAmWithoutTax>1037.74</TotalAmWithoutTax>
  <TotalTaxAm>62.26</TotalTaxAm>
  <SellerName>武汉某某酒店管理有限公司</SellerName>
  <BuyerName>某某公司</BuyerName>
  <Remark>恩施建始龙坪网格10千伏线路 1815P825000P</Remark>
</EInvoice>
"""

# 配置里记着一个「老目录」（3 个文件），界面这次要扫的是「新测试目录」（1 个文件）。
# 之前后端只读配置、忽略界面传的路径，于是清空台账想重新测试时，进来的永远是老目录那堆。
stale = TMP / "老目录"
stale.mkdir(exist_ok=True)
for i in range(3):
    (stale / f"老{i}.pdf").write_bytes(b"%PDF-1.4\n%%EOF\n")
app_web.save_config({"source": str(stale), "recursive": True})

fresh = TMP / "新测试目录"
fresh.mkdir(exist_ok=True)
(fresh / "dzfp_26349119423005870096.xml").write_text(INV_XML, encoding="utf-8")

api3 = app_web.Api()
api3._import_worker(True, "", "", "", str(fresh))       # 同步跑，不起线程、不弹窗
res = api3._result
chk(res.get("files") == 1,
    f"扫的是界面传的目录（1 个文件），不是配置里的老目录（3 个）→ 实际 {res.get('files')}")
chk(Path(str(res.get("source", ""))).name == "新测试目录",
    f"入库结果里写明了本次扫描目录（{res.get('source')}）")
saved = json.loads(app_paths.CONFIG_FILE.read_text(encoding="utf-8")).get("source")
chk(saved == str(fresh), f"扫完顺手把界面路径记进配置（{saved}）")
chk(res.get("added", 0) >= 1 and len(ledger.load()) >= 1,
    f"新目录里的凭证确实进了台账（added={res.get('added')}，台账 {len(ledger.load())} 行）")

# 传空 source 时仍然退回配置里的目录（老行为保留，避免界面漏传就报错）
app_web.save_config({"source": str(stale), "recursive": True})
api4 = app_web.Api()
api4._import_worker(True, "", "", "", "")
chk(api4._result.get("files") == 3, f"不传 source 时退回配置目录（3 个）→ {api4._result.get('files')}")

api.clear_ledger()

print("\n" + "=" * 66)
if FAIL:
    print(f"❌ 有 {len(FAIL)} 项没通过：")
    for x in FAIL:
        print("   -", x)
    sys.exit(1)
print("✅ 全部通过")
print(f"（临时目录：{TMP}）")
