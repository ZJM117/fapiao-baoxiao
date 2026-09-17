# -*- coding: utf-8 -*-
"""
隔离自检：① 打车行程单跟在发票后面（附件归位）② 报销单里的差旅费补助
========================================================================
对应用户 2026-09-15 的两条要求：
  「打车行程单要跟在那个发票之后，然后合成 pdf 的时候也要和进去，因为他是证明
    这个发票的附件要在一块，你可以考虑金额什么的」
  「在生成报销单的时候，要能够增加差旅费，因为出差有差旅费啊，要计算几个人几天，
    多少钱，可以我自己设置这个，我自己填写」

验证点：
  A. 配对（按金额 + 同文件夹/同人加分，一对一）
  B. 排序（行程单紧跟发票，没配上的留在原地）
  C. 金额口径（附件不重复计钱；张数照实算）
  D. 差旅费补助（人数×天数×标准；手改金额优先；空行不污染单据）
  E. 出单据（PDF / Excel 里真有「差旅费补助」和正确的合计）
  F. 合并（发票+行程单落在同一张 A4 的上下两半）

全部在 %TEMP% 副本里跑，绝不碰真实台账 / 输出 / 配置。无窗口（只有无头 Edge 造测试 PDF）。
跑法：  python 测试\\diag_attach_allow.py
"""
import io
import sys
import tempfile
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

TMP = Path(tempfile.mkdtemp(prefix="inv_att_"))

# ⚠️ 必须在 import ledger / app_web / report_pdf 之前改路径
import app_paths                                                    # noqa: E402
app_paths.LEDGER_FILE = TMP / "发票台账.xlsx"
app_paths.LEDGER_BAK = TMP / "台账备份"
app_paths.OUTPUT_DIR = TMP / "输出"
app_paths.WORK_DIR = TMP / "缓存"
app_paths.CONFIG_FILE = TMP / "config.json"
app_paths.LAUNCH_LOG = TMP / "launch.log"
app_paths.UI_PORT_FILE = TMP / "ui_port.txt"
for _d in (app_paths.LEDGER_BAK, app_paths.OUTPUT_DIR, app_paths.WORK_DIR):
    _d.mkdir(parents=True, exist_ok=True)

import attachment as at                                             # noqa: E402
import app_web                                                      # noqa: E402
import ledger                                                       # noqa: E402
import report_pdf as rp                                             # noqa: E402

FAIL = []


def chk(cond, label, extra=""):
    print(("  [OK] " if cond else "  [XX] ") + label + ("" if cond else f"   → {extra}"))
    if not cond:
        FAIL.append(label)


def section(t):
    print("\n" + "=" * 70 + f"\n{t}\n" + "=" * 70)


# ---------------------------------------------------------------- 造假数据
SRC = TMP / "报销"
FOLDERS = {k: SRC / k for k in ("张三", "李四", "王五", "赵六", "钱七")}
for d in FOLDERS.values():
    d.mkdir(parents=True, exist_ok=True)


def mkpdf(folder, name, tag):
    """造一张可读的假票 PDF（票面里放个独一无二的标记，方便验证「和谁在同一页」）"""
    p = FOLDERS[folder] / name
    rp.html_to_pdf(f"""<!doctype html><html><head><meta charset="utf-8"><style>
    @page {{ size: A4; margin: 14mm; }} body {{ font: 13px/1.8 sans-serif; }}
    </style></head><body><h1>电子凭证</h1><div>{tag}</div></body></html>""", p)
    return p


P = {
    "invA": mkpdf("张三", "某某科技-烟台站打车发票76.10.pdf", "TAG-INV-A 打车电子发票 76.10"),
    "tripA": mkpdf("张三", "某某科技-烟台站行程报销单.pdf", "TAG-TRIP-A 行程单 76.10"),
    "invA2": mkpdf("李四", "别的打车发票76.10.pdf", "TAG-INV-A2 另一张 76.10"),
    "invB": mkpdf("王五", "合肥南-财智中心13.58.pdf", "TAG-INV-B 打车发票 13.58"),
    "tripB": mkpdf("王五", "合肥南-财智中心行程单.pdf", "TAG-TRIP-B 行程单 13.58"),
    "tripC": mkpdf("赵六", "没有对应发票的行程单9.99.pdf", "TAG-TRIP-C 行程单 9.99"),
    "train": mkpdf("钱七", "武汉-宜昌东263.pdf", "TAG-TRAIN 火车票 263.00"),
}


def row(seq, ctype, no, date, total, path, person, **kw):
    r = {"序号": seq, "状态": "未报销", "凭证类型": ctype, "发票号码": no,
         "开票日期": date, "票种": "", "行程/明细": kw.get("trip", ""), "购方名称": "某某公司",
         "销方名称": kw.get("seller", ""), "不含税金额": "", "税额": "", "价税合计": total,
         "项目/事由": kw.get("proj", ""), "项目编号": "", "费用类别": kw.get("cat", "交通费"),
         "报销人": person, "报销批次": "", "入库时间": "2026-09-15 15:00",
         "文件路径": str(path), "提示": ""}
    return r


ROWS = [
    row(1, "发票", "26317000003130255066", "2026-08-31", 76.10, P["invA"], "张三",
        proj="某某科技-烟台站", seller="某某汽车服务有限公司"),
    row(2, "发票", "26317000003130255067", "2026-08-31", 76.10, P["invA2"], "李四",
        proj="另一趟打车", seller="某某出租"),
    row(3, "打车行程单", "", "2026-09-01", 76.10, P["tripA"], "张三",
        trip="悦行优选  某某科技→烟台站  08-31 18:09  1 笔行程  76.10元", seller="高德地图"),
    row(4, "发票", "26332000007500040846", "2026-09-02", 13.58, P["invB"], "王五",
        proj="合肥南-财智中心", seller="某某科技"),
    row(5, "打车行程单", "", "2026-09-02", 13.58, P["tripB"], "王五",
        trip="添猫出行  合肥南→财智中心  09-02 11:11  1 笔行程  13.58元", seller="高德地图"),
    row(6, "打车行程单", "", "2026-09-03", 9.99, P["tripC"], "赵六",
        trip="只有行程单、没有发票  9.99元", seller="高德地图"),
    row(7, "火车票", "26349119423005870096", "2026-09-04", 263.00, P["train"], "钱七",
        trip="G1234 武汉→宜昌东 二等座", seller="中国铁路"),
]
ledger.save(ROWS, do_backup=False)
api = app_web.Api()

# ---------------------------------------------------------------- A. 配对
section("一、附件配对：按金额认，同文件夹/同人的优先")
d0 = at.analyze(ROWS)
pairs = {str(a["序号"]): str(m["序号"]) for a, m, _amt in d0["pairs"]}
print("   配对结果：", pairs)
for a, m, amt in d0["pairs"]:
    print(f"     附件 序号{a['序号']} ¥{amt} → 主票 序号{m['序号']}（{m['项目/事由']}）")
chk(d0["n_attach"] == 2, f"5 张里有 2 张行程单能配上发票（实际 {d0['n_attach']}）", pairs)
chk(pairs.get("3") == "1", "同额的两张发票里，挑了同文件夹 + 同报销人的那张（序号1）", pairs)
chk(pairs.get("5") == "4", "13.58 的行程单配给 13.58 的发票（序号4）", pairs)
chk("6" not in pairs, "没有对应发票的行程单（9.99）不算附件")
chk(at.amount_of(ROWS[2]) == 76.10 and abs(at.amount_of(ROWS[0]) - at.amount_of(ROWS[1])) < 0.01,
    "前提数据无误：两张发票都是 76.10")

section("二、排序：行程单紧跟发票，没配上的留原地")
order = [int(r["序号"]) for r in d0["rows"]]
print("   排好的顺序：", order)
chk(order == [1, 3, 2, 4, 5, 6, 7], "序号 3 紧跟 1、序号 5 紧跟 4，其余不动", order)
chk([int(r["序号"]) for r in ROWS] == [1, 2, 3, 4, 5, 6, 7], "原列表没被就地改乱")

section("三、金额口径：附件不重复计钱")
meta = {"kind": "差旅费报销单", "person": "张三", "dept": "二部", "date": "2026年09月15日",
        "reason": "项目现场审计", "start": "2026-08-31", "end": "2026-09-04",
        "place": "烟台"}
d = rp.report_data(ROWS, meta)
# 凭证 = 76.10(1) + 76.10(2) + 13.58(4) + 9.99(6) + 263(7) = 438.77；附件 76.10+13.58 不计
print(f"   total={d['total']} n={d['n']} n_bill={d['n_bill']} n_attach={d['n_attach']}")
chk(abs(d["total"] - 438.77) < 0.01, f"合计只算凭证 = ¥438.77（实际 {d['total']}）")
chk(d["n"] == 7, f"张数是打印出来的 7 张（实际 {d['n']}）")
chk(d["n_attach"] == 2, f"其中 2 张是附件（实际 {d['n_attach']}）")
chk(abs(sum(v["total"] for v in d["by_cat"].values()) - 438.77) < 0.01,
    "金额汇总里的小计也不含附件", d["by_cat"])
chk(d["by_cat"].get("交通费", {}).get("n") == 7, "但张数把附件也算上（7 张）", d["by_cat"])
chk("证明附件" in d["attach"] and "合计 7 张" in d["attach"], f"附件清单说清了（{d['attach']}）")
chk(str(d["rows"][1]["序号"]) == "3", "单据明细里 序号3 也排在 序号1 后面")

# ---------------------------------------------------------------- D. 差旅费补助
section("四、差旅费补助：人数 × 天数 × 标准，自己填")
d_alw = rp.report_data(ROWS, dict(meta, allowance=[
    {"name": "伙食补助费", "people": 3, "days": 4, "rate": 100},        # 自动算 1200
    {"name": "市内交通费", "people": 2, "days": 3, "rate": 80},          # 自动算 480
    {"name": "杂费", "people": 1, "days": 1, "rate": 0, "amount": 156.5},  # 手改金额优先
    {"name": "", "people": "", "days": "", "rate": ""},                  # 空行丢掉
]))
print("   补助明细：", d_alw["alw"])
chk(len(d_alw["alw"]) == 3, f"3 项有效补助（空行丢掉）（实际 {len(d_alw['alw'])}）", d_alw["alw"])
chk(abs(d_alw["alw"][0]["amount"] - 1200.00) < 0.01, "3 人 × 4 天 × 100 = 1,200.00",
    d_alw["alw"][0])
chk(abs(d_alw["alw"][1]["amount"] - 480.00) < 0.01, "2 人 × 3 天 × 80 = 480.00", d_alw["alw"][1])
chk(abs(d_alw["alw"][2]["amount"] - 156.50) < 0.01, "手填的金额优先（156.50）", d_alw["alw"][2])
chk(abs(d_alw["alw_total"] - 1836.50) < 0.01, f"补助合计 = 1,836.50（实际 {d_alw['alw_total']}）")
chk(abs(d_alw["grand"] - (438.77 + 1836.50)) < 0.01,
    f"报销合计 = 凭证 438.77 + 补助 1836.50 = 2,275.27（实际 {d_alw['grand']}）")
chk(d_alw["upper"] == rp.rmb_upper(d_alw["grand"]) and "贰仟贰佰" in d_alw["upper"],
    f"大写跟着报销合计走（{d_alw['upper']}）")
chk(abs(d["grand"] - d["total"]) < 0.01, "没填补助时，报销合计就是凭证合计（回归）")
chk(d_alw["days"] == "5 天", f"出差天数仍按起止算（实际 {d_alw['days']}）")

# ---------------------------------------------------------------- E. 出单据
section("五、生成 PDF + Excel：单据里真有差旅费补助")
res = api.make_report(dict(meta, allowance=[
    {"name": "伙食补助费", "people": 3, "days": 4, "rate": 100},
    {"name": "市内交通费", "people": 2, "days": 3, "rate": 80},
    {"name": "杂费", "people": 1, "days": 1, "rate": 0, "amount": 156.5},
]), [r["序号"] for r in ROWS], "both")
chk(not res.get("error"), "make_report 成功（PDF + Excel）", res.get("error"))
if not res.get("error"):
    print("   ", {k: res.get(k) for k in ("count", "bill", "alw", "sum")})
    chk(abs(res["sum"] - 2275.27) < 0.01, f"返回的报销合计 = 2,275.27（实际 {res['sum']}）")
    chk(abs(res["bill"] - 438.77) < 0.01, f"返回的凭证金额 = 438.77（实际 {res['bill']}）")
    chk(res.get("attach") == 2, f"返回的附件数 = 2（实际 {res.get('attach')}）")
    pdf = [p for p in res["outs"] if p.endswith(".pdf")][0]
    xlsx = [p for p in res["outs"] if p.endswith(".xlsx")][0]
    from pypdf import PdfReader                                      # noqa: E402
    txt = "".join((p.extract_text() or "") for p in PdfReader(pdf).pages)
    flat = "".join(txt.split())
    for want in ("差旅费报销单", "差旅费补助", "补助项目", "人数", "天数",
                 "标准（元/人·天）", "伙食补助费", "市内交通费", "差旅费补助合计",
                 "凭证金额", "报销合计", "2,275.27", "1,836.50", rp.rmb_upper(2275.27),
                 "打车行程单（附件）", "证明附件"):
        chk(want.replace(" ", "") in flat, f"PDF 里有「{want}」")
    chk("438.77" in txt, "PDF 里凭证金额 438.77 也在（2,275.27 - 1,836.50 对得上）")
    import openpyxl                                                   # noqa: E402
    ws = openpyxl.load_workbook(xlsx).worksheets[0]
    cells = [str(c.value) for r in ws.iter_rows() for c in r if c.value is not None]
    xflat = " | ".join(cells)
    for want in ("差旅费报销单", "差旅费补助（3 项）", "差旅费补助合计", "伙食补助费",
                 "市内交通费", "杂费", "报销合计", "凭证金额", rp.rmb_upper(2275.27)):
        chk(want in xflat, f"Excel 里有「{want}」")
    money_cells = [c for r in ws.iter_rows() for c in r
                   if isinstance(c.value, (int, float)) and abs(float(c.value) - 2275.27) < 0.005]
    chk(money_cells and all("#,##0.00" in c.number_format for c in money_cells),
        "Excel 里报销合计是数值 + 千分位", [c.number_format for c in money_cells])
    alw_cells = [c for r in ws.iter_rows() for c in r
                 if isinstance(c.value, (int, float)) and abs(float(c.value) - 1200.0) < 0.005]
    chk(alw_cells and all("#,##0.00" in c.number_format for c in alw_cells),
        "Excel 里补助金额也是数值 + 千分位")

# ---------------------------------------------------------------- F. 合并
section("六、合并 PDF：发票和它的行程单落在同一张 A4 上")
res2 = api.merge_invoices([r["序号"] for r in ROWS], "2up")
print("   ", {k: res2.get(k) for k in ("files", "pages", "asked", "linked", "missing")})
chk(not res2.get("error"), "合并成功", res2.get("error"))
chk(res2.get("files") == 7, f"7 张全合并进去（实际 {res2.get('files')}）")
chk(res2.get("linked") == 2, f"报告了 2 组「发票+行程单」（实际 {res2.get('linked')}）")
chk(res2.get("pages") == 4, f"2up 版式下 4 页 A4（实际 {res2.get('pages')}）")
from pypdf import PdfReader as _R                                     # noqa: E402
pages = [p.extract_text() or "" for p in _R(res2["out"]).pages]
for i, t in enumerate(pages):
    print(f"   第 {i + 1} 页：{' / '.join(x.strip() for x in t.split(chr(10)) if x.strip())[:90]}")
chk("TAG-INV-A" in pages[0] and "TAG-TRIP-A" in pages[0],
    "① 第一页 = 打车发票 + 它的行程单（上下两张）", pages[0][:80].replace("\n", " "))
chk("TAG-INV-B" in pages[2] and "TAG-TRIP-B" in pages[2],
    "② 另一组也在同一页（发票 + 行程单）", pages[2][:80].replace("\n", " "))
chk("TAG-INV-A2" not in pages[0], "同额的那张别的发票没被错配进来")

section("七、合并（每页一张）也按「发票→行程单」的顺序")
res3 = api.merge_invoices([r["序号"] for r in ROWS], "plain")
t3 = [_R(res3["out"]).pages[i].extract_text() or "" for i in range(len(_R(res3["out"]).pages))]
chk(res3.get("pages") == 7, f"每页一张 → 7 页（实际 {res3.get('pages')}）")
chk(t3[0].find("TAG-INV-A") >= 0 and t3[1].find("TAG-TRIP-A") >= 0,
    "第 1 页发票、第 2 页就是它的行程单")
chk(res3.get("linked") == 2, f"plain 也报告了 2 组（实际 {res3.get('linked')}）")

# ---------------------------------------------------------------- G. 界面数据
section("八、界面：台账顺序 + 「↳ 附件」标记")
led = api.get_ledger({})
chk([int(r["序号"]) for r in led["rows"]] == [1, 3, 2, 4, 5, 6, 7],
    "台账返回的顺序里行程单跟在发票后面", [r["序号"] for r in led["rows"]])
chk(set(led.get("attach_links", {}).keys()) == {"3", "5"},
    f"带回了附件标记（实际 {list(led.get('attach_links', {}).keys())}）",
    led.get("attach_links"))
chk(led["attach_links"].get("3", {}).get("main_seq") == "1",
    "标记里写明是第 1 号发票的附件", led["attach_links"].get("3"))
chk(led["attach_links"].get("5", {}).get("main_amount") == 13.58,
    "标记里带上了主发票金额（13.58）", led["attach_links"].get("5"))
# 2026-09-17 起台账页「合计」改成跟单据同一个口径（附件不重复计）：
# 全量数字挪到 sum_raw 留作对照，被排除的附件金额单独给 attach_sum。
chk(abs(led["sum"] - 438.77) < 0.01,
    f"台账「合计」＝计费口径 438.77（行程单 76.10+13.58 已含在发票里，不再加）实际 {led['sum']}")
chk(abs(led.get("sum_raw", 0) - 528.45) < 0.01,
    f"全部行硬加 sum_raw 仍是全量 528.45（实际 {led.get('sum_raw')}）")
chk(abs(led.get("attach_sum", 0) - 89.68) < 0.01,
    f"被排除的附件金额 = 76.10+13.58 = 89.68（实际 {led.get('attach_sum')}）")

section("九、筛选后只选一张时，别把它当附件（避免金额对不上）")
led2 = api.get_ledger({"ctype": "打车行程单"})
chk([int(r["序号"]) for r in led2["rows"]] == [3, 5, 6], "只筛行程单 → 3 条",
    [r["序号"] for r in led2["rows"]])
chk(not led2.get("attach_links"), "拉不到发票时就不标记为附件（后端会照常计钱）",
    led2.get("attach_links"))
d2 = rp.report_data([r for r in ROWS if r["凭证类型"] == "打车行程单"], {"kind": "费用报销单"})
chk(abs(d2["total"] - 99.67) < 0.01,
    f"只选 3 张行程单时全部计钱 = 76.10+13.58+9.99 = 99.67（实际 {d2['total']}）")

# ---------------------------------------------------------------- 收尾
print("\n" + "=" * 70)
if FAIL:
    print(f"❌ 有 {len(FAIL)} 项没过：")
    for f in FAIL:
        print("   ·", f)
    sys.exit(1)
print("✅ 全部通过")
