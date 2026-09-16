# -*- coding: utf-8 -*-
"""
模板二版式（出差报销明细）自检 —— 全程不碰真实台账
====================================================
用户 2026-09-15：「报销单还需要一个别的模板，这个设置为模板二。
                   用这个模板做成这个的」（参考 E:\\桌\\报销明细-某某公司-7.28.xlsx）

这个检查盯的是：
  · 版式骨架（标题 / 抬头行 / 时间+同行人 / 明细六列 / 小计 / 按人补助 / 合计）有没有走对
  · 按人汇总的口径：票据合计只算报销凭证，附件（打车行程单）不重复计钱
  · 「类型」列的短标签映射（高铁票 / 打车 / 住宿 / 机票 / 公交 …）
  · 票面标记（退票 / 改签）有没有写进「备注」
  · Excel 和 PDF 出自同一份行模型（行数、数字必须一模一样）
  · 一分补助都没填时，不出「出差补助」那一块

跑法：
  C:\\Users\\JM\\.workbuddy\\binaries\\python\\envs\\gui\\Scripts\\python.exe 测试\\diag_jy.py
退出码 0 = 全通过。产物写在 %TEMP%\\jy_diag\\，不落进同步盘、不碰真实台账。
"""
import os
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl                       # noqa: E402
import report_pdf as rp               # noqa: E402

fails = []


def chk(cond, label, extra=""):
    print(("  ok   " if cond else "  FAIL ") + label + ("" if cond else "  <<< " + str(extra)))
    if not cond:
        fails.append(label)


OUT = Path(tempfile.gettempdir()) / "jy_diag"
OUT.mkdir(parents=True, exist_ok=True)


def R(seq, ctype, no, date, trip, total, who="", mark="", cat="交通费", proj="", folder="批次A"):
    """造一条真实形状的台账行"""
    return {
        "序号": str(seq), "状态": "未报销", "凭证类型": ctype, "票面标记": mark,
        "出行人": who, "发票号码": no, "开票日期": date, "票种": "",
        "行程/明细": trip, "购方名称": "国网湖北省电力有限公司", "销方名称": "某某公司",
        "不含税金额": "", "税额": "", "价税合计": total, "项目/事由": proj,
        "项目编号": "", "费用类别": cat, "报销人": "张三", "报销批次": "",
        "入库时间": "2026-09-15 10:00",
        "文件路径": rf"E:\桌\样本\{folder}\dzfp_{no or seq}.pdf", "提示": "",
    }


ROWS = [
    R(1, "火车票", "", "2026-08-17", "G1234  济南东-烟台  2026-08-17 08:08  二等座",
      174.00, "李强"),
    # 下面这两条是一对：金额分毫不差 → 打车行程单是这张发票的证明附件，不重复计钱
    R(2, "发票", "26349119423005870001", "2026-08-18", "", 76.10, "李强"),
    R(3, "打车行程单", "", "2026-08-18", "济南西-大明湖", 76.10, "李强"),
    R(4, "发票", "26349119423005870002", "2026-08-20", "", 208.00, "李贵清", cat="住宿费"),
    R(5, "火车票", "", "2026-08-21", "K123  烟台-济南  2026-08-21 09:30  硬座",
      50.00, "李贵清", mark="退票"),
    R(6, "飞机行程单", "", "2026-08-22", "武汉-烟台  李贵清", 748.00, "李贵清"),
    R(7, "打车行程单", "", "2026-08-22", "地铁2号线 五四广场-青岛站", 12.00, "", folder="批次B"),
    R(8, "发票", "26349119423005870003", "2026-08-23", "", 30.00, "", cat="办公费"),
]

META = {
    "kind": rp.JY_KIND, "person": "", "dept": "",
    "reason": "国网湖北决算审核，第五批29个项目，审计现场进点差旅费",
    "date": "2026年09月15日", "start": "2026-08-17", "end": "2026-08-23",
    "allowance_by": [{"name": "李强", "amount": 1200}, {"name": "李贵清", "amount": 600}],
}

print("— 版式判定 —")
chk(rp.JY_KIND == "模板二", "单据类型名就是用户指定的那一串")
chk(rp.JY_TITLE == "出差报销明细", "标题照参考文件叫「出差报销明细」")
chk(rp.style_of(rp.JY_KIND) == "jy", "选它 → 走 jy 版式")
chk(rp.style_of("差旅费报销单") == "trip", "差旅费报销单还是老版式")
chk(rp.style_of("费用报销单") == "cost" and rp.style_of("") == "cost",
    "费用报销单（含空值）还是老版式")
chk(rp.style_of(" 模板二 ") == "jy", "两头有空格也认得出来")

print("\n— 类型列的短标签（照模板的 高铁票 / 打车 / 住宿 / 公交）—")
chk(rp._short_type(ROWS[0]) == "高铁票", "火车票 G 字头 → 高铁票")
chk(rp._short_type(ROWS[4]) == "火车票", "火车票 K 字头 → 火车票")
chk(rp._short_type(ROWS[2]) == "打车", "打车行程单 → 打车")
chk(rp._short_type(ROWS[6]) == "公交", "行程里写地铁 → 公交")
chk(rp._short_type(ROWS[3]) == "住宿", "发票 + 费用类别住宿费 → 住宿")
chk(rp._short_type(ROWS[5]) == "机票", "飞机行程单 → 机票")
chk(rp._short_type(ROWS[7]) == "办公", "发票 + 办公费 → 办公")

print("\n— 按人汇总的口径（附件不重复计钱）—")
d = rp.report_data(ROWS, META)
chk(d["style"] == "jy" and d["title"] == "出差报销明细", "report_data 认得这是 jy 版式")
chk(d["n_attach"] == 1, f"认出 1 组附件（实际 {d['n_attach']}）")
# 报销凭证：174 + 76.10 + 208 + 50 + 748 + 12 + 30 = 1298.10（那条 76.10 的行程单是附件，不算）
chk(abs(d["total"] - 1298.10) < 0.001, f"票据合计 1298.10（实际 {d['total']}）")
by = {p["name"]: p for p in d["by_person"]}
chk(abs(by["李强"]["bills"] - 250.10) < 0.001,
    f"李强的票据合计 250.10＝174＋76.10（实际 {by['李强']['bills']}）")
chk(abs(by["李贵清"]["bills"] - 1006.00) < 0.001,
    f"李贵清的票据合计 1006.00＝208＋50＋748（实际 {by['李贵清']['bills']}）")
chk("（未标注）" in by, "没认出行人的票归到「（未标注）」，不会被漏掉")
chk(abs(by["（未标注）"]["bills"] - 42.00) < 0.001, "（未标注）＝12＋30")
chk(abs(sum(p["bills"] for p in d["by_person"]) - d["total"]) < 0.001,
    "按人加起来 ＝ 票据合计（一个人都没漏）")
chk(abs(by["李强"]["alw"] - 1200) < 0.001 and abs(by["李强"]["real"] - 1450.10) < 0.001,
    "李强：实际＝票据 250.10＋补助 1200＝1450.10")
chk(abs(by["（未标注）"]["alw"]) < 0.001, "没填补助的人，补助为 0")
chk(abs(d["alw_total"] - 1800) < 0.001, "补助合计 ＝ 1200＋600")
chk(abs(d["grand"] - 3098.10) < 0.001, f"合计＝票据 1298.10＋补助 1800＝3098.10（实际 {d['grand']}）")
chk(d["person"] == "李强、李贵清",
    f"抬头「报销人」多人时自动把名单填上（实际 {d['person']!r}）")
chk(d["travelers"] == ["李强", "李贵清"], f"同行人名单（实际 {d['travelers']}）")
chk(d["alw"] == [], "jy 版式不走「人数×天数×标准」那块（不会两套补助一起算）")

print("\n— 一分补助都没填 → 不出补助块 —")
d0 = rp.report_data(ROWS, {**META, "allowance_by": []})
chk(d0["by_person"] and d0["alw_total"] == 0, "人还在名单里，只是补助为 0")
_, lay0 = rp._jy_layout(ROWS, {**META, "allowance_by": []})
flat0 = [cell[0] for row in lay0 for cell in row]
chk("出差补助" not in flat0, "行模型里不出现「出差补助」那一块")
chk("合计" in flat0, "「合计」照旧在")

print("\n— 行模型（PDF / Excel 共用这一份）—")
_, lay = rp._jy_layout(ROWS, META)
vals = [c[0][0] for c in lay]
chk(vals[0] == "出差报销明细" and lay[0][0][1] == 6, "第 1 行是标题，横跨 6 列")
chk(vals[1] == "报销人" and lay[1][1][0] == "李强、李贵清"
    and lay[1][3][0].startswith("国网湖北决算审核"),
    "第 2 行是「报销人｜姓名｜日期｜事由｜报销金额｜备注」", [c[0] for c in lay[1]])
chk(lay[2][0][0] == "时间：2026-08-17 - 2026-08-23" and lay[2][0][1] == 2,
    "第 3 行左边「时间」并两格", lay[2])
chk(lay[2][2][0] == "同行人：李强、李贵清" and lay[2][2][1] == 3,
    "右边「同行人」并三格（2+1+3＝6 列不多不少）", lay[2])
n_detail = len(ROWS)
chk(len(lay) == 3 + n_detail + 1 + 1 + 3 + 1,
    f"行数 = 抬头3＋明细{n_detail}＋小计1＋补助表头1＋3人＋合计1（实际 {len(lay)}）")
chk(vals[3 + n_detail] == "小计" and abs(lay[3 + n_detail][3][0] - d["total"]) < 0.001,
    "明细后跟「小计」（A:B 合并的标签 + E 列的票据合计）", lay[3 + n_detail])
chk(vals[-1] == "合计" and abs(lay[-1][4][0] - d["grand"]) < 0.001, "最后一行「合计」＝票据＋补助")
# 附件那一行：备注写「（附件）」，金额照旧显示（方便核对）
def cell(row, i):
    """取某一行的第 i 格（行长短不一，越界给 None）"""
    return row[i][0] if i < len(row) else None


att_row = [r for r in lay if cell(r, 5) == "（附件）"]
chk(len(att_row) == 1 and abs(cell(att_row[0], 4) - 76.10) < 0.001,
    "附件行标了「（附件）」，金额照旧摆出来")
mark_row = [r for r in lay if cell(r, 5) == "退票"]
chk(len(mark_row) == 1, "票面标记（退票）写进「备注」列")
chk([cell(lay[3], i) for i in (0, 1, 2)] == ["李强", "高铁票", "2026-08-17"],
    "明细列序：出行人｜类型｜日期")
chk(str(cell(lay[3], 3)).startswith("G1234"), "第 4 格是行程 / 事由")
chk(cell(lay[3], 5) == "", "没标记没附件的行，备注列是空的")

print("\n— PDF（无头 Edge 打印）—")
outs = rp.make_report(ROWS, META, str(OUT / "jy样例"), fmt="both")
pdf = [p for p in outs if str(p).lower().endswith(".pdf")]
xls = [p for p in outs if str(p).lower().endswith(".xlsx")]
chk(len(pdf) == 1 and Path(pdf[0]).exists() and Path(pdf[0]).stat().st_size > 3000,
    "PDF 出得来", pdf)
if pdf:
    head = Path(pdf[0]).read_bytes()[:5]
    chk(head == b"%PDF-", "是正经 PDF 文件")
html = rp.build_jy_html(ROWS, META)
chk("出差报销明细" in html and 'colspan="6"' in html, "HTML 里有标题和六列骨架")
chk("李强、李贵清" in html and "报销人" in html, "HTML 抬头里有报销人和姓名")
chk("时间：2026-08-17 - 2026-08-23" in html and "同行人：李强、李贵清" in html,
    "HTML 里有时间 / 同行人")
chk("1,298.10" in html and "3,098.10" in html, "HTML 里小计 / 合计都对")
chk("（附件）" in html and "退票" in html, "HTML 里附件与票面标记都在备注列")
chk("出差补助" in html and "李贵清" in html, "HTML 里有按人的补助块")

print("\n— 类型跟附件走 / 列宽 / 尾块不分页 —")
# 打车数电票本身是张「发票」、费用类别也就写「交通费」，光看它会映射成「交通」，
# 可它那张打车行程单写的是「打车」—— 同一笔钱两行两个叫法。配上的附件是什么类型，
# 主票就跟着写什么（走 attachment 的 link_of）。
types_by_seq = {str(r.get("序号")): cell(lay[3 + i], 1) for i, r in enumerate(d["rows"])}
chk(types_by_seq["2"] == "打车",
    f"配了打车行程单的发票，类型也写「打车」（实际 {types_by_seq.get('2')!r}）")
chk(types_by_seq["3"] == "打车", "行程单自己还是「打车」")
chk(types_by_seq["4"] == "住宿" and types_by_seq["6"] == "机票",
    "住宿费发票→住宿、飞机行程单→机票")
chk(types_by_seq["7"] == "公交", "行程里带「地铁」的 → 公交")
chk(types_by_seq["5"] == "火车票", "K 字头普速 → 火车票（不是高铁票）")

# 列宽：HTML 之前没给 colgroup，table-layout:fixed 把 6 列平摊，
# 「行程」被挤成窄条、地址换行 4 行，白多出一页。
pcts = rp._jy_col_pct()
chk(len(pcts) == 6 and abs(sum(pcts) - 100) < 0.5, f"六列百分比加起来 ≈100（实际 {sum(pcts)}）")
chk(pcts[3] == max(pcts) and abs(pcts[3] - 45.8 / sum(rp._JY_COLS) * 100) < 0.1,
    "第 4 列（行程）最宽，跟比较文件的列宽比例一致")
chk(html.count("<col ") == 6 and f'<col style="width:{pcts[0]}%">' in html,
    "HTML 的表头里带 6 个列宽")

# 分页：小计 → 补助块 → 合计 必须抱在一起（30 行真实台账曾被拆成两页、
# 末页只剩一行「合计」）
tail = rp._jy_tail_index(d)
chk(cell(lay[tail], 0) == "小计" and cell(lay[-1], 0) == "合计",
    "尾块正好从「小计」起、到「合计」止")
chk(html.count("<tbody") == 2 and '<tbody class="keep">' in html,
    "HTML 把尾块单独包了一层 tbody.keep")
chk("tbody.keep" in rp._JY_CSS and "page-break-inside: avoid" in rp._JY_CSS,
    "CSS 里有「整块不许跨页」的规则")
chk(html.index('<tbody class="keep">') > html.index("G1234"),
    "尾块那一层排在明细后面（顺序没被打乱）")

print("\n— Excel —")
chk(len(xls) == 1 and Path(xls[0]).exists(), "Excel 出得来", xls)
if xls:
    wb = openpyxl.load_workbook(xls[0])
    ws = wb["报销单"]
    chk(ws.title == "报销单", "工作表就叫「报销单」（跟参考文件一致）")
    chk(ws["A1"].value == "出差报销明细", "A1 是标题")
    chk("A1:F1" in [str(r) for r in ws.merged_cells.ranges], "标题横跨 A1:F1")
    chk(ws["A2"].value == "报销人" and str(ws["B2"].value or "").startswith("李强"),
        "第 2 行抬头（报销人 + 姓名）")
    chk(ws["A3"].value.startswith("时间：") and str(ws["D3"].value or "").startswith("同行人："),
        "第 3 行时间 / 同行人")
    chk("A3:B3" in [str(r) for r in ws.merged_cells.ranges]
        and "D3:F3" in [str(r) for r in ws.merged_cells.ranges],
        "第 3 行的两处合并跟参考文件一样（A:B / D:F）")
    # 明细从第 4 行开始
    chk(ws["A4"].value == "李强" and ws["B4"].value == "高铁票" and ws["C4"].value == "2026-08-17",
        "明细第 1 行：出行人 / 类型 / 日期")
    chk(ws["D4"].value.startswith("G1234"), "明细第 1 行第 4 格是行程")
    money_cell = ws["E4"]
    chk(isinstance(money_cell.value, (int, float)) and abs(money_cell.value - 174) < 0.001,
        "金额写的是**数字**（不是字符串），能直接再算", repr(money_cell.value))
    chk(money_cell.number_format == "#,##0.00", "金额套了千分位两位小数格式")
    chk(ws["A1"].font.name == "微软雅黑" and ws["A1"].font.size == 14 and ws["A1"].font.bold,
        "标题字体照参考文件（微软雅黑 14 加粗）")
    chk(ws["A4"].font.name == "微软雅黑" and ws["A4"].font.size == 11, "正文微软雅黑 11")
    chk(ws["A2"].border.left.style == "thin" and ws["F2"].border.bottom.style == "thin",
        "有细边框（参考文件是满框线的）")
    chk(ws.column_dimensions["D"].width > 40, f"「行程」列够宽（{ws.column_dimensions['D'].width}）")
    chk(ws.row_dimensions[1].height == 33 and ws.row_dimensions[2].height == 60.6,
        "标题 / 抬头行的行高照参考文件")
    # 小计 / 合计 / 补助块
    small = 3 + n_detail + 1
    chk(ws.cell(row=small, column=1).value == "小计"
        and abs(ws.cell(row=small, column=5).value - d["total"]) < 0.001,
        "小计行位置和金额都对")
    chk(ws.cell(row=small + 1, column=5).value == "出差补助"
        and ws.cell(row=small + 1, column=6).value == "实际",
        "补助块表头照参考文件放在 E/F 两格")
    chk(ws.cell(row=small + 2, column=2).value == "出差补助"
        and ws.cell(row=small + 2, column=4).value == "李强"
        and abs(ws.cell(row=small + 2, column=5).value - 1200) < 0.001
        and abs(ws.cell(row=small + 2, column=6).value - 1450.10) < 0.001,
        "补助每人一行：B列标签 / D列姓名 / E列补助 / F列实际",
        [ws.cell(row=small + 2, column=c).value for c in range(1, 7)])
    last = small + 2 + 3
    chk(ws.cell(row=last, column=1).value == "合计"
        and abs(ws.cell(row=last, column=5).value - d["grand"]) < 0.001, "合计行")
    chk(f"$F${last + 2}" in str(ws.print_area),
        f"打印区域连页脚说明一起算（A1:F{last + 2}）", ws.print_area)
    chk(ws.page_setup.orientation == "portrait" and ws.page_setup.fitToWidth == 1,
        "A4 竖版、按宽度缩到一页")
    chk(ws.freeze_panes == "A3", "冻结在 A3（跟参考文件一样）")
    # 页脚那几行不带边框，别把说明画成表格
    foot = last + 1
    chk(ws.cell(row=foot, column=1).border.left.style is None,
        "页脚说明没有边框（不跟表格混在一起）")
    chk("制表时间" in str(ws.cell(row=foot, column=1).value or "")
        or any("制表时间" in str(ws.cell(row=last + i, column=1).value or "") for i in (1, 2)),
        "页脚写了制表时间")

print("\n— 不破坏老版式 —")
old = rp.report_data(ROWS, {"kind": "差旅费报销单", "person": "李强",
                            "allowance": [{"name": "伙食补助费", "people": 3, "days": 4, "rate": 100}]})
chk(old["style"] == "trip" and abs(old["alw_total"] - 1200) < 0.001,
    "差旅费报销单的补助还算人数×天数×标准")
chk(rp.build_report_html(ROWS, {"kind": "费用报销单"}).count("<h1>") == 1,
    "费用报销单还是老 HTML（标题走 h1）")
_, lay_old = rp._jy_layout(ROWS, {"kind": "费用报销单"})
chk(len(lay_old) > 3, "行模型函数本身不挑单据类型（只认 kind 里的名字）")

print()
if fails:
    print(f"❌ {len(fails)} 项不通过：" + "；".join(fails[:6]))
    sys.exit(1)
print("✅ 全部通过")
