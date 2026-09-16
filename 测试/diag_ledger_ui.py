# -*- coding: utf-8 -*-
"""凭证/台账/后端接口 的隔离端到端自检
=====================================
在 %TEMP% 里跑完整链路：解析真实票据文件夹 → 入库临时台账 → 再调 app_web 的后端接口，
验证「凭证类型 / 行程明细」两列、按类型筛选、按类型统计这些新加的东西真的通了。
不碰真实台账、不写真实输出、不弹任何窗口。

跑法：
  C:\\Users\\JM\\.workbuddy\\binaries\\python\\envs\\gui\\Scripts\\python.exe 测试\\diag_ledger_ui.py
"""
import shutil
import sys
import tempfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SRC = Path(r"E:\桌\湖北报销发票整理")
tmp = Path(tempfile.mkdtemp(prefix="invui_"))

import app_paths  # noqa: E402
app_paths.LEDGER_FILE = tmp / "发票台账.xlsx"
app_paths.LEDGER_BAK = tmp / "台账备份"
app_paths.WORK_DIR = tmp / "缓存"
app_paths.OUTPUT_DIR = tmp / "输出"
app_paths.CONFIG_FILE = tmp / "config.json"
app_paths.LAUNCH_LOG = tmp / "launch.log"
app_paths.UI_PORT_FILE = tmp / "ui_port.txt"
for _d in (app_paths.LEDGER_BAK, app_paths.WORK_DIR, app_paths.OUTPUT_DIR):
    _d.mkdir(parents=True, exist_ok=True)

import invparse  # noqa: E402
import ledger  # noqa: E402
import app_web  # noqa: E402

fail = []


def check(name, ok, extra=""):
    print(f"  {'✅' if ok else '❌'} {name}" + (f"　{extra}" if extra else ""))
    if not ok:
        fail.append(name)


try:
    print(f"数据源：{SRC}（存在={SRC.exists()}）")

    # ---------- 1. 解析 ----------
    files = invparse.collect_files(SRC)
    recs = invparse.dedupe_same_invoice([r for r in (invparse.parse_file(f) for f in files) if r])
    print(f"\n扫描 {len(files)} 个文件 → 去重后 {len(recs)} 条凭证")

    # ---------- 2. 入库 ----------
    res = ledger.add_records(recs, batch="2026-09 某某公司")
    added = len(res["added"])
    check("入库行数 > 80", added > 80, f"{added} 条")
    check("需人工 ≤ 2", len(res["manual"]) <= 2, f"{len(res['manual'])} 个")

    # 后端 _import_worker 里那段「按凭证类型拆分」的逻辑，这里原样跑一遍
    by_type = {}
    for r in res["added"]:
        t = str(r.get("凭证类型") or "其他凭证")
        d = by_type.setdefault(t, {"n": 0, "sum": 0.0})
        d["n"] += 1
        d["sum"] += float(str(r.get("价税合计") or 0).replace(",", "") or 0)
    print("  本次新增构成：" + "　".join(f"{t} {v['n']}条/¥{v['sum']:,.2f}"
                                    for t, v in sorted(by_type.items(), key=lambda kv: -kv[1]["n"])))
    check("按凭证类型拆分出 ≥3 类", len(by_type) >= 3, "、".join(by_type))

    # ---------- 3. 台账两列真的写进去了 ----------
    rows = ledger.load()
    check("台账落盘行数一致", len(rows) == added, f"{len(rows)} 行")
    check("每一行都有「凭证类型」", all(str(r.get("凭证类型") or "").strip() for r in rows),
          "空的：" + str(sum(1 for r in rows if not str(r.get("凭证类型") or "").strip())))
    has_detail = [r for r in rows if str(r.get("行程/明细") or "").strip()]
    check("有「行程/明细」的行 > 50", len(has_detail) > 50, f"{len(has_detail)} 行")
    cnt = Counter(str(r.get("凭证类型")) for r in rows)
    print("  凭证类型分布：" + "、".join(f"{k} {v}" for k, v in cnt.most_common()))

    # ---------- 4. 后端接口 ----------
    api = app_web.Api()
    r0 = api.get_ledger({})
    check("get_ledger 返回 rows", len(r0.get("rows") or []) == len(rows), f"{len(r0.get('rows') or [])}")
    check("get_ledger 带 ctypes 下拉项", bool(r0.get("ctypes")), str(r0.get("ctypes")))
    check("aggregates 含「按凭证类型」", "按凭证类型" in (r0.get("aggregates") or {}),
          str(list((r0.get("aggregates") or {}).keys())))
    check("get_ledger 每行都带新列",
          all("凭证类型" in x and "行程/明细" in x for x in r0["rows"]))

    top_type = cnt.most_common(1)[0][0]
    r1 = api.get_ledger({"ctype": top_type})
    check(f"按凭证类型筛选（{top_type}）", len(r1["rows"]) == cnt[top_type],
          f"{len(r1['rows'])} vs {cnt[top_type]}")

    # 关键词：行程里才有的词（车次/机场/起终点）
    sample = next((r for r in rows if str(r.get("行程/明细") or "").strip()), None)
    if sample:
        word = str(sample["行程/明细"]).split()[0][:6]
        r2 = api.get_ledger({"kw": word})
        check(f"关键词能搜到行程明细（{word}）", any(
            str(x["序号"]) == str(sample["序号"]) for x in r2["rows"]), f"命中 {len(r2['rows'])} 行")

    # 关键词：项目 / 编号
    proj = next((r for r in rows if str(r.get("项目编号") or "").strip()), None)
    if proj:
        r3 = api.get_ledger({"kw": str(proj["项目编号"])})
        check(f"关键词能搜到项目编号（{proj['项目编号']}）", len(r3["rows"]) > 0, f"命中 {len(r3['rows'])} 行")

    # 行程单没有开票日期时，按月筛应该退回入库时间，不能筛成空
    nodate = [r for r in rows if not str(r.get("开票日期") or "").strip()]
    if nodate:
        mon = str(nodate[0].get("入库时间") or "")[:7]
        r4 = api.get_ledger({"month": mon})
        check(f"无开票日期的行按入库月筛（{mon}）不为空", len(r4["rows"]) > 0, f"命中 {len(r4['rows'])} 行")

    # ---------- 5. 导出汇总 ----------
    rr = api.export_summary({})
    check("导出汇总成功", bool(rr.get("out")), str(rr))
    import openpyxl
    wb = openpyxl.load_workbook(rr["out"])
    check("汇总 7 张 sheet", len(wb.sheetnames) == 7, "、".join(wb.sheetnames))
    check("汇总含「按凭证类型」sheet", "按凭证类型" in wb.sheetnames)

    # ---------- 6. 二次入库不新增 ----------
    res2 = ledger.add_records(
        invparse.dedupe_same_invoice([r for r in (invparse.parse_file(f) for f in files) if r]))
    check("第二次入库新增 0 行", len(res2["added"]) == 0, f"实际 {len(res2['added'])}")

    print()
    print(f"❌ 未通过 {len(fail)} 项：" + "、".join(fail) if fail else "✅ 全部通过")
finally:
    shutil.rmtree(tmp, ignore_errors=True)

sys.exit(1 if fail else 0)
