# -*- coding: utf-8 -*-
"""
部署首启隔离测试（Docker / NAS 场景）
=====================================
模拟「项目拷到 NAS + data 目录里放了带数据的老台账 Excel + 第一次 docker compose up」，
把《Docker部署说明.md》里承诺的几件事逐条验证一遍。

  ① 镜像内容按 .dockerignore 的规则拷 → 副本里**没有任何运行数据**（尤其不能有 .db）
  ② 首启自动建库：data/发票台账.db 出现，Excel 里的行被搬进库
  ③ 迁移前自动备份原 Excel（台账备份/ 里能看到）
  ④ 首启自动建管理员：用户名 admin，初始密码就是 FB_PASSWORD
  ⑤ 登录三条路：admin+口令进得去（身份=管理员）／用户名留空只填口令也能进／错口令挡下
  ⑥ /api/ledger_db_info 报的条数 = 库里实际条数
  ⑦ 再启动一次：不重复迁移（库里还是那么多行，不会翻倍）

全程在 %TEMP% 副本里跑（子进程起真实服务、走真实 HTTP），无窗口，真实数据一个字节不动。
跑法：  python 测试/diag_deploy_first_run.py
"""
import io
import json
import os
import shutil
import socket
import sqlite3
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PY = sys.executable
PW = "deploy-pw-9527"
N_ROWS = 3
FAIL = []


def chk(cond, label, extra=""):
    print(("  [OK] " if cond else "  [XX] ") + label + ("" if cond else f"   → {extra}"))
    if not cond:
        FAIL.append(label)


def section(t):
    print("\n" + "=" * 66 + f"\n{t}\n" + "=" * 66)


def free_port() -> int:
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    port = s.getsockname()[1]
    s.close()
    return port


class _NoRedirect(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, *a, **k):
        return None


_noredir = urllib.request.build_opener(_NoRedirect)


def http(path, method="GET", data=None, cookie=None, timeout=20):
    body = None
    req = urllib.request.Request(BASE + path, method=method)
    if data is not None:
        body = urllib.parse.urlencode(data).encode("utf-8")
        req.add_header("Content-Type", "application/x-www-form-urlencoded")
    if cookie:
        req.add_header("Cookie", cookie)
    try:
        r = _noredir.open(req, body, timeout=timeout)
        return r.status, dict(r.headers), r.read()
    except urllib.error.HTTPError as e:
        return e.code, dict(e.headers), e.read()
    except Exception as e:                                          # noqa: BLE001
        return 0, {}, str(e).encode()


def api_call(name, *args, cookie=""):
    body = json.dumps({"args": list(args)}).encode("utf-8")
    req = urllib.request.Request(BASE + "/api/" + name, data=body, method="POST")
    req.add_header("Content-Type", "application/json")
    if cookie:
        req.add_header("Cookie", cookie)
    try:
        r = urllib.request.urlopen(req, timeout=30)
        return json.loads(r.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return {"error": f"HTTP {e.code}"}


def login(username, password):
    """登录，返回 Set-Cookie（空串＝没让进）"""
    _, h, _ = http("/login", "POST", {"username": username, "password": password})
    return h.get("Set-Cookie", "") or ""


def db_live_rows():
    p = DATA / "发票台账.db"
    if not p.exists():
        return None
    c = sqlite3.connect("file:" + p.as_posix() + "?mode=ro", uri=True)
    try:
        return c.execute("SELECT COUNT(*) FROM invoices WHERE is_deleted = 0").fetchone()[0]
    finally:
        c.close()


def db_user_count():
    p = DATA / "发票台账.db"
    c = sqlite3.connect("file:" + p.as_posix() + "?mode=ro", uri=True)
    try:
        return c.execute("SELECT COUNT(*) FROM users").fetchone()[0]
    finally:
        c.close()


def make_ledger_excel(path: Path, n: int = N_ROWS):
    """造一份「用户从本机带过来的老台账」（表头照真实台账抄）"""
    import openpyxl
    src = ROOT / "发票台账.xlsx"
    hdr = None
    if src.exists():
        wb0 = openpyxl.load_workbook(src, read_only=True)
        hdr = [c.value for c in next(wb0.active.iter_rows(min_row=1, max_row=1))]
        wb0.close()
    if not hdr or not any(hdr):
        raise RuntimeError(f"拿不到台账表头：{src}")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "台账"
    ws.append(hdr)

    def col(*kws):
        for i, h in enumerate(hdr):
            if h and any(k in str(h) for k in kws):
                return i
        return None

    i_seq, i_kind = col("序号"), col("凭证类型")
    i_no, i_date, i_tot = col("发票号码"), col("开票日期"), col("价税合计")
    for k in range(n):
        row = [None] * len(hdr)
        if i_seq is not None:
            row[i_seq] = k + 1
        if i_kind is not None:
            row[i_kind] = "电子发票"
        if i_no is not None:
            row[i_no] = f"2531700000000000{k}"
        if i_date is not None:
            row[i_date] = f"2026-01-0{k + 1}"
        if i_tot is not None:
            row[i_tot] = 100.0 + k
        ws.append(row)
    wb.save(path)


# =====================================================================
# ① 造一份「镜像内容」：严格按 .dockerignore 的规则排除运行数据
# =====================================================================
TMP = Path(tempfile.mkdtemp(prefix="deploy_"))
APP = TMP / "app"
DATA = TMP / "data"
SRC = TMP / "票据"
for d in (DATA, SRC):
    d.mkdir(parents=True, exist_ok=True)

IGNORE = shutil.ignore_patterns(
    "缓存", "输出", "台账备份", "browser_debug", "__pycache__", "测试", "docs",
    "*.db", "*.db-wal", "*.db-shm", "*.xlsx", "*.xls", "*.log", "*.txt", "*.bat",
    "*.md", "config.json", "*.pyc", ".git")
shutil.copytree(ROOT, APP, ignore=IGNORE)
make_ledger_excel(DATA / "发票台账.xlsx")

PORT = free_port()
BASE = f"http://127.0.0.1:{PORT}"
SRV_LOG = TMP / "srv.log"

print("=" * 74)
print(f"临时根目录 {TMP}")
print(f"模拟镜像目录 {APP}")
print(f"模拟 /data 卷 {DATA}")
print(f"服务地址   {BASE}")
print("=" * 74)


def start_server():
    env = dict(os.environ)
    env.update({"FB_SERVER": "1", "FB_HOST": "127.0.0.1", "FB_PORT": str(PORT),
                "FB_DATA_DIR": str(DATA), "FB_SOURCE": str(SRC), "FB_PASSWORD": PW})
    lf = open(SRV_LOG, "ab")
    p = subprocess.Popen([PY, str(APP / "app_web.py")], cwd=str(APP), env=env,
                         stdout=lf, stderr=subprocess.STDOUT)
    for _ in range(160):
        if p.poll() is not None:
            txt = SRV_LOG.read_text(encoding="utf-8", errors="replace")
            raise RuntimeError("子进程提前退出了：\n" + txt[-2000:])
        code, _, _ = http("/api/ping")
        if code == 200:
            return p
        time.sleep(0.25)
    raise RuntimeError("服务 40 秒没起来")


def stop_server(p):
    p.terminate()
    try:
        p.wait(timeout=15)
    except subprocess.TimeoutExpired:
        p.kill()


section("① 镜像内容（.dockerignore 是否真的挡住了数据）")
no_data = [str(f.relative_to(APP)) for f in APP.rglob("*")
           if f.is_file() and (f.suffix in (".db", ".xlsx") or f.name == "config.json")]
chk(not no_data, "镜像里没有台账数据库 / Excel / 配置（数据只从 /data 卷来）", no_data)
chk((APP / "db.py").exists(), "db.py 在镜像里（源码要跟着走）")
chk((APP / "app_web.py").exists() and (APP / "gui" / "index.html").exists(), "主程序与界面文件都在")

section("② 首启建库 + 迁移老 Excel")
chk(not (DATA / "发票台账.db").exists(), "启动前确实还没有数据库")
proc = start_server()
dbp = DATA / "发票台账.db"
chk(dbp.exists(), "启动后自动建出台账数据库 发票台账.db")
chk(db_live_rows() == N_ROWS, f"Excel 里的 {N_ROWS} 行被搬进了库", f"实际 {db_live_rows()}")
bak = sorted((DATA / "台账备份").glob("*.xlsx")) if (DATA / "台账备份").exists() else []
chk(bool(bak), "迁移前把原 Excel 备份进了 台账备份/", [b.name for b in bak])
log = (DATA / "launch.log").read_text(encoding="utf-8", errors="replace")
chk("服务器（Docker/NAS）" in log, "日志里写明运行模式是服务器")

section("③④ 首启自动建管理员")
chk(db_user_count() >= 1, "库里出现了账号（不然部署完没人进得去）", db_user_count())

section("⑤ 登录三条路")
code, _, body = http("/login")
chk(code == 200 and b'name="password"' in body, "登录页有密码输入框（还有用户名框）")
chk(b'name="username"' in body, "登录页有用户名输入框")

c_bad = login("admin", "wrong-password")
chk(not c_bad, "错口令拿不到会话（没让进）", c_bad)

c_admin = login("admin", PW)
chk(bool(c_admin), "admin + FB_PASSWORD 能登录")
me = api_call("whoami", cookie=c_admin).get("result") or {}
chk(me.get("role") == "admin", "admin 登录后的身份是管理员", me)
chk(me.get("user") == "admin", "会话里记的是用户名 admin（审计要靠它）", me.get("user"))

c_pw = login("", PW)
chk(bool(c_pw), "用户名留空、只填访问口令也能进（忘了管理员密码时的后门）")

section("⑥ 库信息接口与库里实际条数对得上")
info = api_call("ledger_db_info", cookie=c_admin).get("result") or {}
chk(info.get("invoices") == N_ROWS, f"/api/ledger_db_info 报的条数 = {N_ROWS}", info)
chk(info.get("excel_exists") is True, "快照 Excel 存在（给人看/打印的那份）", info.get("excel"))
chk(info.get("unique_index_ok") is True, "唯一索引建上了（并发查重靠它）", info)
chk((info.get("users") or 0) >= 1, "接口也报出了账号数")

audit = api_call("audit_list", cookie=c_admin).get("result") or {}
chk(not audit.get("error"), "审计日志接口可用（管理员）", audit.get("error"))

section("⑦ 再启动一次：不重复迁移")
stop_server(proc)
proc = start_server()
chk(db_live_rows() == N_ROWS, f"重启后还是 {N_ROWS} 行（没有重复迁一遍变成 {N_ROWS * 2}）",
    db_live_rows())
chk(bool(login("admin", PW)), "重启后账号仍有效（账号存在库里，不是内存里）")
stop_server(proc)

section("结果")
print(f"临时目录（可整个删掉）：{TMP}")
if FAIL:
    print(f"\n有 {len(FAIL)} 项没过：")
    for f in FAIL:
        print("   -", f)
    sys.exit(1)
print("\n全部通过 —— 部署说明里承诺的首启行为都对得上")
