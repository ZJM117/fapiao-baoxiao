# -*- coding: utf-8 -*-
"""台账隔离测试：在 %TEMP% 里跑一遍「解析 → 入库 → 再入库查重」。
不碰真实台账、不弹窗。
"""
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SRC = Path(r"E:\桌\湖北报销发票整理")
tmp = Path(tempfile.mkdtemp(prefix="invled_"))

import app_paths  # noqa: E402
app_paths.LEDGER_FILE = tmp / "发票台账.xlsx"
app_paths.LEDGER_BAK = tmp / "台账备份"
app_paths.WORK_DIR = tmp / "工作"
app_paths.OUTPUT_DIR = tmp / "输出"

import invparse  # noqa: E402
import ledger  # noqa: E402

fail = []


def check(name, ok, extra=""):
    print(f"  {'✅' if ok else '❌'} {name}" + (f"　{extra}" if extra else ""))
    if not ok:
        fail.append(name)


try:
    files = invparse.collect_files(SRC)
    print(f"扫描到 {len(files)} 个文件")
    recs = [invparse.parse_file(f) for f in files]
    recs = [r for r in recs if r]
    recs = invparse.dedupe_same_invoice(recs)
    print(f"去重后 {len(recs)} 条凭证\n")

    print("=== 第一次入库 ===")
    res = ledger.add_records(recs, person="", batch="2026-09 某某公司")
    n_added = len(res["added"])
    check("新增行数 > 80", n_added > 80, f"{n_added} 行")
    check("需人工只有极少数（图片/损坏件）", len(res["manual"]) <= 2,
          f"{len(res['manual'])} 个：" + "、".join(
              Path(m.get("src", "")).name[:24] for m in res["manual"]))
    rows = ledger.load()
    check("台账落盘", app_paths.LEDGER_FILE.exists() and len(rows) == n_added,
          f"{len(rows)} 行")

    # 按凭证类型统计
    from collections import Counter
    cnt = Counter(str(r.get("凭证类型", "")) for r in rows)
    print("     凭证类型分布：" + "、".join(f"{k or '(空)'} {v} 条" for k, v in cnt.most_common()))

    # 关键字段完整度
    miss_total = [r for r in rows if not r.get("价税合计")]
    miss_date = [r for r in rows if not r.get("开票日期")]
    check("缺金额的行 ≤ 2", len(miss_total) <= 2, f"缺 {len(miss_total)} 行")
    for r in miss_total:
        print(f"        · 无金额：[{r.get('凭证类型')}] {Path(str(r.get('文件路径'))).name[:46]}")
    check("缺日期的行 ≤ 3", len(miss_date) <= 3, f"缺 {len(miss_date)} 行")
    for r in miss_date:
        print(f"        · 无日期：[{r.get('凭证类型')}] {Path(str(r.get('文件路径'))).name[:46]}")
    has_detail = [r for r in rows if r.get("行程/明细")]
    check("有行程明细的行 > 60", len(has_detail) > 60, f"{len(has_detail)} 行")

    # 抽查几条
    print("\n     抽查（各类型第一条）：")
    seen = set()
    for r in rows:
        t = str(r.get("凭证类型", ""))
        if t in seen:
            continue
        seen.add(t)
        print(f"       [{t}] {str(r.get('发票号码','')) or '—':22s} "
              f"{str(r.get('开票日期','')):11s} {str(r.get('价税合计','')):>9s} "
              f"{(r.get('费用类别') or ''):6s} {(r.get('行程/明细') or '')[:40]}")

    print("\n=== 第二次入库（同一批，应该全部认出来、不新增）===")
    recs2 = [invparse.parse_file(f) for f in files]
    recs2 = invparse.dedupe_same_invoice([r for r in recs2 if r])
    res2 = ledger.add_records(recs2, batch="2026-09 某某公司")
    check("第二次新增 0 行", len(res2["added"]) == 0, f"实际新增 {len(res2['added'])} 行")
    for r in res2["added"]:
        print(f"        · 被当成新记录：[{r.get('凭证类型')}] {r.get('发票号码') or '—'} "
              f"{r.get('开票日期')} {r.get('价税合计')} {str(r.get('行程/明细'))[:36]}")
    check("第二次全部认出来", len(res2["in_ledger"]) + len(res2["dup"]) >= len(recs2) - 1,
          f"已有 {len(res2['in_ledger'])}，重复风险 {len(res2['dup'])}")

    print("\n=== 标已报销 + 再扫一次：应该报重复 ===")
    seqs = [r["序号"] for r in rows[:20]]
    n = ledger.update_rows(seqs, 状态="已报销")
    check("批量标记 20 行", n == 20, f"实际 {n} 行")
    res3 = ledger.add_records(recs2, batch="2026-09 某某公司")
    check("标已报销的再入库会报重复风险", len(res3["dup"]) > 0, f"{len(res3['dup'])} 条")

    print("\n=== 汇总 ===")
    a = ledger.aggregate(ledger.load())
    for k in ("按凭证类型", "按费用类别", "按状态"):
        items = "、".join(f"{kk} {vv['张数']}张/{vv['价税合计']:.2f}"
                         for kk, vv in sorted(a[k].items(), key=lambda x: -x[1]["价税合计"])[:5])
        print(f"     {k}：{items}")

    out = tmp / "汇总.xlsx"
    ledger.export_summary(out, ledger.load())
    import openpyxl
    wb = openpyxl.load_workbook(str(out))
    check("导出汇总含 7 张 sheet", len(wb.sheetnames) == 7, "、".join(wb.sheetnames))

    print()
    if fail:
        print(f"❌ 未通过 {len(fail)} 项：" + "、".join(fail))
    else:
        print("✅ 全部通过")
finally:
    shutil.rmtree(tmp, ignore_errors=True)

sys.exit(1 if fail else 0)
