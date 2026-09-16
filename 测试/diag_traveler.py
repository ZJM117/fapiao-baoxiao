# -*- coding: utf-8 -*-
"""
隔离自检：出行人认领（三层兜底 + 手机号自动学习 + 人工手填）
================================================================================
用户提的问题：「要是提取人名好提取吗，因为有的行程单没有人名什么的，就是打车的单子」

事实（用真实票据核对过）：
  · 火车票 → 票面有乘车人姓名；飞机行程单 → 票面有旅客姓名：能直接读。
  · 打车**发票**（数电票）→ 「出行人」列有时填（赵大勇那张）、有时空着。
  · 打车**行程单** → 票面**完全没有姓名**，只有「行程人手机号」。
    所以要看三层兜底：票面 → 手机号对照表 → 文件夹名；都没认出来就留空让人手填。

全部在 %TEMP% 副本里做，真实台账一个字都不写；真实票据只读。
跑法：  python 测试\\diag_traveler.py
"""
import io
import re
import shutil
import sys
import tempfile
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
BASE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BASE))

TMP = Path(tempfile.mkdtemp(prefix="inv_tv_"))

# ⚠️ 必须在 import ledger / app_web 之前改路径（它们用 from app_paths import XXX 取值）
import app_paths                                                    # noqa: E402
app_paths.LEDGER_FILE = TMP / "发票台账.xlsx"
app_paths.LEDGER_BAK = TMP / "台账备份"
app_paths.OUTPUT_DIR = TMP / "输出"
app_paths.WORK_DIR = TMP / "缓存"
app_paths.CONFIG_FILE = TMP / "config.json"
app_paths.LAUNCH_LOG = TMP / "launch.log"
app_paths.UI_PORT_FILE = TMP / "ui_port.txt"
for _d in (app_paths.LEDGER_BAK, app_paths.OUTPUT_DIR, app_paths.WORK_DIR):
    _d.mkdir(parents=True, exist_ok=True)

import app_web                                                      # noqa: E402
import invparse as ip                                               # noqa: E402
import ledger as lg                                                 # noqa: E402
import traveler as tv                                               # noqa: E402

FAIL = []


def chk(cond, label, extra=""):
    print(("  [OK] " if cond else "  [XX] ") + label + ("" if cond else f"   → {extra!r}"))
    if not cond:
        FAIL.append(label)


def rec(ctype="", traveler="", tv_from="", phone="", src="", date="", trip_date="",
        total=None, no="", warn=None):
    r = ip._blank(src, "test", ctype)
    r.update({"traveler": traveler, "tv_from": tv_from, "phone": phone,
              "date": date, "trip_date": trip_date, "total": total, "no": no})
    if warn is not None:
        r["warn"] = list(warn)
    return r


def fresh_map():
    """每次用例前清空对照表，免得互相影响"""
    p = tv.map_path()
    if p.exists():
        p.unlink()
    return {}


# =====================================================================
print("\n=== 1. 什么像个人名 ===")
for good in ("王小明", "李强", "陈晓燕", "赵大勇", "欧阳娜娜", "尉迟小花"):
    chk(tv.is_name(good), f"「{good}」是个人名")
for bad in ("", "13800000000", "模板二", "某某科技", "某某科技",
            "旅客运输服务", "出租车", "行程单", "A123", "李强2", "某某公司"):
    chk(not tv.is_name(bad), f"「{bad}」不算人名")

# =====================================================================
print("\n=== 2. 文件夹名认人（第三层兜底） ===")
cases = [("E:/a/王小明（某某公司）/x.pdf", "王小明"),
         ("E:/a/李强（某某公司）/x.pdf", "李强"),
         ("E:/a/陈晓燕（某某公司）/x.pdf", "陈晓燕"),
         ("E:/a/赵大勇/x.pdf", "赵大勇")]
for p, want in cases:
    got = tv.from_folder(p)
    chk(got == want, f"{Path(p).parent.name} → {want}", got)
for p in ("E:/a/报销/x.pdf", "E:/a/新建文件夹/x.pdf", "E:/a/2026-09/x.pdf",
          "E:/a/待报销发票/x.pdf", ""):
    chk(tv.from_folder(p) == "", f"「{Path(p).parent.name if p else '（空路径）'}」不会被人名", tv.from_folder(p))

# =====================================================================
print("\n=== 3. 三层顺序：票面 → 手机号 → 文件夹 ===")
fresh_map()
P = "E:/报销/王小明（某某公司）/"

# ① 票面有名字 → 用票面，不看后面的
r = rec(traveler="李强", tv_from="票面", phone="13800000000", src=P + "a.pdf")
n, s = tv.resolve(r)
chk((n, s) == ("李强", tv.SRC_FACE), "票面有名字就先信票面", (n, s))

# ② 老台账里被误存成一串手机号的，不认，往下走
r = rec(traveler="13800000000", tv_from="", phone="13800000000", src=P + "a.pdf")
n, s = tv.resolve(r)
chk(n == "王小明" and s == tv.SRC_FOLDER, "旧的手机号脏值不会被当成人名，继续走兜底", (n, s))

# ③ 有对照表 → 用手机号翻人名（打车行程单主要靠这层）
tv.set_name("13800000000", "王小明")
r = rec(phone="13800000000", src=P + "a.pdf")
n, s = tv.resolve(r)
chk((n, s) == ("王小明", tv.SRC_PHONE), "手机号对照表命中 → 出行人来自「手机号」", (n, s))

# ④ 对照表里没这个号 → 退到文件夹
r = rec(phone="13900000000", src=P + "a.pdf")
n, s = tv.resolve(r)
chk((n, s) == ("王小明", tv.SRC_FOLDER), "号没记过 → 退到文件夹兜底", (n, s))

# ⑤ 三层都没有 → 留空（交给人工填），绝不瞎猜
r = rec(src="E:/报销/2026-09/a.pdf")
n, s = tv.resolve(r)
chk((n, s) == ("", ""), "三层都没线索 → 留空，不猜", (n, s))

# ⑥ 人工改过的对照表要能改、能删
tv.set_name("13800000000", "陆小玲")
chk(tv.load_map().get("13800000000") == "陆小玲", "对照表能人工改名")
tv.set_name("13800000000", "")
chk("13800000000" not in tv.load_map(), "名字给空 = 删掉这条对照")

# =====================================================================
print("\n=== 4. 自动学习（手机号 → 姓名） ===")
fresh_map()
F1 = "E:/报销/王小明（某某公司）/"
F2 = "E:/报销/李强（某某公司）/"
batch = [
    # 同文件夹的火车票：票面有姓名（学习素材）
    rec("火车票", traveler="王小明", tv_from="票面", src=F1 + "t1.pdf", date="2026-08-31"),
    # 同一次出差的行程单：只有手机号，日期接近 → 应该学会
    rec("打车行程单", phone="13800000000", src=F1 + "p1.pdf", trip_date="2026-08-31"),
    # 日期差太远 → 不学（不硬凑）
    rec("打车行程单", phone="13700000001", src=F1 + "p2.pdf", trip_date="2025-01-01"),
    # 别的文件夹的票，不该被拿来当证据
    rec("打车行程单", phone="13600000002", src=F2 + "p3.pdf", trip_date="2026-08-31"),
    # 文件直接摊在根目录 → 不学
    rec("打车行程单", phone="13500000003", src="E:/报销/p4.pdf", trip_date="2026-08-31"),
]
stat = tv.apply(batch)
L = stat["learned"]
chk(L == {"13800000000": "王小明"}, "只学了「同文件夹 + 日期接近」的那条", L)
chk(tv.load_map().get("13800000000") == "王小明", "学到的写进了对照表")
chk(len(tv.load_map()) == 1, "没学到的没被写进去", tv.load_map())
chk(stat[tv.SRC_FACE] == 1 and stat[tv.SRC_PHONE] == 1 and stat[tv.SRC_FOLDER] == 2,
    "各来源统计对得上（票面1 / 手机号1 / 文件夹2 / 没线索1）", stat)
chk(stat[""] == 1, "散在根目录、三层都没线索的那张留空（不瞎猜）", stat)

# 同一文件夹里有两个人 → 谁都别学（宁可让人手填）
fresh_map()
amb = [
    rec("火车票", traveler="王小明", tv_from="票面", src=F1 + "t1.pdf", date="2026-08-31"),
    rec("火车票", traveler="李强", tv_from="票面", src=F1 + "t2.pdf", date="2026-08-31"),
    rec("打车行程单", phone="13911112222", src=F1 + "p1.pdf", trip_date="2026-08-31"),
]
tv.apply(amb)
chk("13911112222" not in tv.load_map(), "同文件夹有两个候选人 → 不猜、不学", tv.load_map())

# 已经记过的号，学到新证据也不覆盖（以人工/最早的为准）
fresh_map()
tv.set_name("13800000000", "王小明")
again = [
    rec("火车票", traveler="别人", tv_from="票面", src=F1 + "t1.pdf", date="2026-08-31"),
    rec("打车行程单", phone="13800000000", src=F1 + "p1.pdf", trip_date="2026-08-31"),
]
tv.apply(again)
chk(tv.load_map()["13800000000"] == "王小明", "已有对照不会被自动学覆盖", tv.load_map())

# 学完以后，同一部手机下次直接认出来
fresh_map()
tv.set_name("13800000000", "王小明")
r = rec("打车行程单", phone="13800000000", src="E:/别的地方/随便/a.pdf")
n, s = tv.resolve(r)
chk((n, s) == ("王小明", tv.SRC_PHONE), "换个文件夹、同一部手机也能认出来（学习的价值）", (n, s))

# =====================================================================
print("\n=== 5. 解析层：行程单的手机号不再当人名 ===")
SRC = Path(r"E:\桌\湖北报销发票整理\报销")
F_TRIP = SRC / "王小明（某某公司）" / "【合肥南-财智中心-13.58元-1个行程】高德打车电子行程单.pdf"
F_CAR_INV = SRC / "赵大勇（某某公司）" / "某某科技公司-烟台站打车票电子发票76.1.pdf"
F_TRIP2 = SRC / "赵大勇（某某公司）" / "某某科技公司-烟台站行程报销单.pdf"
F_TRAIN = SRC / "李强（某某公司）" / "济南东-烟台 李强 10.pdf"

if F_TRIP.exists():
    r = ip.parse_file(F_TRIP)
    chk(r.get("phone") == "13800000000", "行程单：手机号进 phone 字段", r.get("phone"))
    chk(r.get("traveler") == "", "行程单：手机号不再被塞进「出行人」", r.get("traveler"))
else:
    print("  (跳过：真实行程单不在)")

if F_TRIP2.exists():
    r = ip.parse_file(F_TRIP2)
    chk(r.get("phone") == "13900000000", "悦行优选行程单也读到了手机号", r.get("phone"))
    chk(r.get("traveler") == "", "悦行优选行程单出行人同样留空", r.get("traveler"))

if F_CAR_INV.exists():
    r = ip.parse_file(F_CAR_INV)
    chk(r.get("traveler") == "赵大勇", "打车发票票面上的「出行人」读到了（原来漏读）", r.get("traveler"))
    chk(r.get("tv_from") == "票面", "来源标成票面", r.get("tv_from"))
    chk("某某科技" in str(r.get("route")), "折行的出发地拼回来了", r.get("route"))
else:
    print("  (跳过：真实打车发票不在)")

if F_TRAIN.exists():
    r = ip.parse_file(F_TRAIN)
    chk(r.get("traveler") == "李强" and r.get("tv_from") == "票面",
        "火车票乘车人照旧读得到", (r.get("traveler"), r.get("tv_from")))

# =====================================================================
print("\n=== 6. 台账「出行人」列 + 重新识别不覆盖手填 ===")
cols = [c[0] for c in lg.COLUMNS]
chk("出行人" in cols, "台账里有「出行人」这一列")
chk(cols.index("出行人") == cols.index("票面标记") + 1, "它就排在票面标记后面",
    f"票面标记@{cols.index('票面标记')} 出行人@{cols.index('出行人')}")
chk("出行人" in lg.FIELDS, "出行人可被人工修改（在白名单里）")

api = app_web.Api()
fresh_map()
rows = [rec("火车票", traveler="王小明", tv_from="票面",
            src="E:/报销/王小明（某某公司）/t.pdf", date="2026-08-31", no="",
            total=204.0, warn=[])]
rows[0]["no"] = "26379217601000118606"
lg.save(lg.add_records(rows)["added"] + [], do_backup=False) if False else None
lg.add_records(rows)
saved = lg.load()
chk(len(saved) == 1 and saved[0].get("出行人") == "王小明",
    "入库后台账「出行人」写上了姓名", [r.get("出行人") for r in saved])
chk(saved[0].get("报销人") == "", "「报销人」不再被出行人顶替（两个概念分开）", saved[0].get("报销人"))

# 人工把出行人改成别的名字，再「重新识别」——不能被冲掉
lg.update_rows([str(saved[0]["序号"])], **{"出行人": "手动改的名字"})
n = api.refresh_rows([str(saved[0]["序号"])])
after = lg.load()
chk(n.get("updated") is not None, "重新识别跑通了", n)
chk(after[0].get("出行人") == "手动改的名字", "手填的出行人不会被重新识别覆盖（原件不在也不清空）",
    after[0].get("出行人"))

# =====================================================================
print("\n=== 7. 报销单上也看得到出行人（明细列 + 出差人兜底） ===")
import report_pdf as rp                                             # noqa: E402

rr = [
    {"序号": "1", "凭证类型": "火车票", "票面标记": "", "开票日期": "2026-08-31",
     "行程/明细": "G123 济南东-烟台", "费用类别": "交通费", "销方名称": "中国铁路",
     "不含税金额": 199.01, "税额": 4.99, "价税合计": 204.0, "出行人": "王小明"},
    {"序号": "2", "凭证类型": "打车发票", "票面标记": "", "开票日期": "2026-08-30",
     "行程/明细": "合肥南-财智中心", "费用类别": "交通费", "销方名称": "高德",
     "不含税金额": 12.81, "税额": 0.77, "价税合计": 13.58, "出行人": "李强"},
]
d = rp.report_data(rr, {"kind": "差旅费报销单", "start": "2026-08-30", "end": "2026-08-31"})
chk(sorted(d["travelers"]) == ["李强", "王小明"],
    "report_data 里能拿到出行人清单（出行人按行序出现，行序是按日期排的）", d["travelers"])
chk(d["person"] == "李强、王小明", "出差人没填 → 自动用出行人兜底", d["person"])
d2 = rp.report_data(rr, {"kind": "差旅费报销单", "person": "张三"})
chk(d2["person"] == "张三", "出差人自己填了 → 尊重手填，不覆盖", d2["person"])
d3 = rp.report_data(rr, {"kind": "费用报销单", "person": ""})
chk(d3["person"] == "", "费用报销单的「报销人」不做出行人兜底（两个概念分开）", d3["person"])

for kind in ("差旅费报销单", "费用报销单"):
    html = rp.build_report_html(rr, {"kind": kind})
    chk(bool(re.search(r"<th[^>]*>出行人</th>", html)), f"{kind}明细表头里有「出行人」列")
    chk("王小明" in html and "李强" in html, f"{kind}明细里出行人名字进去了")

# Excel 版式自检：有边框的行必须铺满 9 列（合并格算覆盖），防止哪一列漏样式成了「缺口」
import tempfile as _tf                                             # noqa: E402
from openpyxl import load_workbook                                 # noqa: E402
xp = Path(_tf.mkdtemp(prefix="tv_xlsx_")) / "报销单.xlsx"
rp.build_report_xlsx(rr, {"kind": "差旅费报销单"}, xp)
wb = load_workbook(xp)
ws = wb.active
covered = set()
for m in ws.merged_cells.ranges:
    for r_ in range(m.min_row, m.max_row + 1):
        for c_ in range(m.min_col, m.max_col + 1):
            covered.add((r_, c_))


def _has_border(rc):
    cell = ws.cell(row=rc[0], column=rc[1])
    b = cell.border
    return bool(b and (b.left.style or b.right.style or b.top.style or b.bottom.style))


holes = []
for r_ in range(1, ws.max_row + 1):
    cells = [(r_, c_) for c_ in range(1, 10)]
    if any(_has_border(rc) for rc in cells):
        miss = [c_ for c_ in range(1, 10)
                if (r_, c_) not in covered and not _has_border((r_, c_))]
        if miss:
            holes.append((r_, miss))
chk(not holes, "Excel 版式：凡是有边框的行，9 列都铺满（合并格算覆盖，不许有缺口）", holes)
chk(ws.cell(row=8, column=4).value == "出行人", "Excel 明细表第 4 列就是出行人",
    ws.cell(row=8, column=4).value)
shutil.rmtree(xp.parent, ignore_errors=True)

# 末了把隔离目录清掉（里面只有生成的空台账/缓存）
shutil.rmtree(TMP, ignore_errors=True)

print()
if FAIL:
    print(f"❌ {len(FAIL)} 项不通过：")
    for x in FAIL:
        print("   ·", x)
    sys.exit(1)
print("✅ 全部通过")
