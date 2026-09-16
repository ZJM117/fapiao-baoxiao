# -*- coding: utf-8 -*-
"""
SQLite 数据层
=============
为什么要有这一层：原来台账就是一个 Excel，程序「读整张表 → 在内存里改 → 整张写回去」。
单机自用没问题，但只要两个人在同一分钟里各改一次，后保存的那个就会把前一个人的改动
整片覆盖掉 —— 而且是静默的，谁都发现不了。这是多人使用最致命的一条。

所以把「谁是真身」改一下：
    SQLite（发票台账.db）  =  唯一写入源
    发票台账.xlsx          =  导出快照，给人看、给测试看、给 Excel 打开用

任何一次写操作都在一个事务里完成 —— SQLite 的写锁会把并发写串行化，
第二个事务读到的是第一个已提交的结果，于是「同时入库两张票」不会再丢一张。

对外仍然只暴露「中文列名的 dict 列表」，界面上看到的东西一点没变。

表：
    invoices         台账行（软删除，删除可追溯）
    users            用户与角色
    audit_logs       审计日志（谁在什么时候改了什么，改前改后都有）
    jobs             后台任务（长时间操作的任务状态）
    generated_files  生成过的报销单 / 合并文件
    meta             零碎键值（迁移标记、Excel 导出指纹…）
"""
import contextlib
import contextvars
import hashlib
import hmac
import json
import secrets
import sqlite3
import threading
import time
from datetime import datetime
from pathlib import Path

import app_paths


def db_path() -> Path:
    """数据库文件位置 = 台账 Excel 旁边那个同名 .db。

    两件事故意这么做：
    ① 每次现算、不缓存成模块常量 —— 诊断脚本会在 import 之前把 app_paths 里的路径
       改到临时目录（见 测试/ 里那句「必须在 import ledger / app_web 之前改路径」），
       缓存成常量会让那条隔离规则静默失效、直接读写到真实台账上；
    ② 跟着 LEDGER_FILE 走（而不是 DATA_DIR）—— 现有的诊断脚本都只改 LEDGER_FILE，
       跟着它走意味着「台账指到哪，数据库就在哪」，不用去改一堆老脚本。
    """
    return Path(app_paths.LEDGER_FILE).with_suffix(".db")

# ---------------------------------------------------------------- 中英对照
# 台账对外一直用中文列名（界面、Excel、所有测试都按中文键取值），
# 数据库里用英文列名。这张表是两者之间唯一的换算处。
COL2DB = {
    "状态": "status",
    "凭证类型": "ctype",
    "票面标记": "stamp",
    "出行人": "traveler",
    "发票号码": "inv_no",
    "开票日期": "inv_date",
    "票种": "kind",
    "行程/明细": "detail",
    "购方名称": "buyer",
    "销方名称": "seller",
    "不含税金额": "amount",
    "税额": "tax",
    "价税合计": "total",
    "项目/事由": "project",
    "项目编号": "project_code",
    "费用类别": "category",
    "报销人": "person",
    "报销批次": "batch",
    "入库时间": "created_at",
    "文件路径": "src_path",
    "提示": "tip",
}
DB2COL = {v: k for k, v in COL2DB.items()}
MONEY_DB = ("amount", "tax", "total")
# 台账列顺序（序号是算出来的，不存在库里）
FIELDS = (["序号"] + [c for c in
                     ("状态", "凭证类型", "票面标记", "出行人", "发票号码", "开票日期", "票种",
                      "行程/明细", "购方名称", "销方名称", "不含税金额", "税额", "价税合计",
                      "项目/事由", "项目编号", "费用类别", "报销人", "报销批次", "入库时间",
                      "文件路径", "提示")])

# ---------------------------------------------------------------- 上下文
# 「这次操作是谁、从哪个 IP 来的」—— ledger 层要写审计，但它不知道 HTTP 的事，
# 所以由 app_web 在请求线程开头 set，ledger 里直接读。
_ctx_user = contextvars.ContextVar("fb_user", default="")
_ctx_ip = contextvars.ContextVar("fb_ip", default="")
_ctx_role = contextvars.ContextVar("fb_role", default="")


def set_actor(user: str = "", ip: str = "", role: str = ""):
    _ctx_user.set(user or "")
    _ctx_ip.set(ip or "")
    _ctx_role.set(role or "")


def actor() -> str:
    return _ctx_user.get() or ""


def role() -> str:
    return _ctx_role.get() or ""


def client_ip() -> str:
    return _ctx_ip.get() or ""


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ======================================================================
# 连接与事务
# ======================================================================
def _open() -> sqlite3.Connection:
    """开一个连接。

    每次操作开新连接（不用线程本地缓存）：http.server 是「一个请求一个新线程」，
    缓存连接反而要操心线程退出时的回收；SQLite 开连接很便宜，WAL 模式下并发读也不阻塞。
    """
    p = db_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(p), timeout=15.0, isolation_level=None)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout = 15000")     # 抢不到写锁就等，别立刻报错
    conn.execute("PRAGMA journal_mode = WAL")       # 读写不互斥（多人同时用才不卡）
    conn.execute("PRAGMA synchronous = NORMAL")
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


@contextlib.contextmanager
def read():
    """只读连接（顺手保证表已经建好）"""
    init()
    conn = _open()
    try:
        yield conn
    finally:
        conn.close()


@contextlib.contextmanager
def write_tx():
    """写事务。

    用 BEGIN IMMEDIATE 一上来就拿写锁 —— 先读后写（DEFERRED）在两个事务同时升级成写
    的时候会互相等，直接 IMMEDIATE 就绕开了这个死锁。拿不到锁会按 busy_timeout 等。
    """
    init()
    conn = _open()
    try:
        conn.execute("BEGIN IMMEDIATE")
        yield conn
        conn.execute("COMMIT")
    except BaseException:
        with contextlib.suppress(Exception):
            conn.execute("ROLLBACK")
        raise
    finally:
        conn.close()


# ======================================================================
# 建表
# ======================================================================
_SCHEMA = """
CREATE TABLE IF NOT EXISTS invoices (
  id           INTEGER PRIMARY KEY AUTOINCREMENT,
  status       TEXT    NOT NULL DEFAULT '未报销',
  ctype        TEXT    NOT NULL DEFAULT '',
  stamp        TEXT    NOT NULL DEFAULT '',
  traveler     TEXT    NOT NULL DEFAULT '',
  inv_no       TEXT    NOT NULL DEFAULT '',
  inv_date     TEXT    NOT NULL DEFAULT '',
  kind         TEXT    NOT NULL DEFAULT '',
  detail       TEXT    NOT NULL DEFAULT '',
  buyer        TEXT    NOT NULL DEFAULT '',
  seller       TEXT    NOT NULL DEFAULT '',
  amount       REAL,
  tax          REAL,
  total        REAL,
  project      TEXT    NOT NULL DEFAULT '',
  project_code TEXT    NOT NULL DEFAULT '',
  category     TEXT    NOT NULL DEFAULT '',
  person       TEXT    NOT NULL DEFAULT '',
  batch        TEXT    NOT NULL DEFAULT '',
  created_at   TEXT    NOT NULL DEFAULT '',
  src_path     TEXT    NOT NULL DEFAULT '',
  tip          TEXT    NOT NULL DEFAULT '',
  fp           TEXT    NOT NULL DEFAULT '',
  path_key     TEXT    NOT NULL DEFAULT '',
  is_deleted   INTEGER NOT NULL DEFAULT 0,
  deleted_at   TEXT    NOT NULL DEFAULT '',
  deleted_by   TEXT    NOT NULL DEFAULT '',
  dupe         INTEGER NOT NULL DEFAULT 0,
  rev          INTEGER NOT NULL DEFAULT 0,
  updated_at   TEXT    NOT NULL DEFAULT '',
  updated_by   TEXT    NOT NULL DEFAULT ''
);
CREATE INDEX IF NOT EXISTS ix_inv_no  ON invoices(inv_no);
CREATE INDEX IF NOT EXISTS ix_fp      ON invoices(fp);
CREATE INDEX IF NOT EXISTS ix_path    ON invoices(path_key);
CREATE INDEX IF NOT EXISTS ix_alive   ON invoices(is_deleted);

CREATE TABLE IF NOT EXISTS users (
  id           INTEGER PRIMARY KEY AUTOINCREMENT,
  username     TEXT    NOT NULL UNIQUE,
  display_name TEXT    NOT NULL DEFAULT '',
  pw_hash      TEXT    NOT NULL,
  role         TEXT    NOT NULL DEFAULT 'biz',
  dept         TEXT    NOT NULL DEFAULT '',
  scope        TEXT    NOT NULL DEFAULT '',
  disabled     INTEGER NOT NULL DEFAULT 0,
  created_at   TEXT    NOT NULL DEFAULT '',
  last_login   TEXT    NOT NULL DEFAULT ''
);

CREATE TABLE IF NOT EXISTS audit_logs (
  id          INTEGER PRIMARY KEY AUTOINCREMENT,
  ts          TEXT    NOT NULL,
  username    TEXT    NOT NULL DEFAULT '',
  ip          TEXT    NOT NULL DEFAULT '',
  action      TEXT    NOT NULL,
  object_type TEXT    NOT NULL DEFAULT '',
  object_id   TEXT    NOT NULL DEFAULT '',
  before_json TEXT,
  after_json  TEXT,
  ok          INTEGER NOT NULL DEFAULT 1,
  error       TEXT    NOT NULL DEFAULT ''
);
CREATE INDEX IF NOT EXISTS ix_audit_ts  ON audit_logs(ts);
CREATE INDEX IF NOT EXISTS ix_audit_obj ON audit_logs(object_type, object_id);
CREATE INDEX IF NOT EXISTS ix_audit_user ON audit_logs(username);

CREATE TABLE IF NOT EXISTS jobs (
  id          TEXT PRIMARY KEY,
  kind        TEXT    NOT NULL DEFAULT '',
  state       TEXT    NOT NULL DEFAULT 'queued',
  total       INTEGER NOT NULL DEFAULT 0,
  done        INTEGER NOT NULL DEFAULT 0,
  label       TEXT    NOT NULL DEFAULT '',
  username    TEXT    NOT NULL DEFAULT '',
  created_at  TEXT    NOT NULL DEFAULT '',
  started_at  TEXT    NOT NULL DEFAULT '',
  ended_at    TEXT    NOT NULL DEFAULT '',
  error       TEXT    NOT NULL DEFAULT '',
  result_json TEXT
);
CREATE INDEX IF NOT EXISTS ix_jobs_state ON jobs(state);

CREATE TABLE IF NOT EXISTS generated_files (
  id            INTEGER PRIMARY KEY AUTOINCREMENT,
  name          TEXT    NOT NULL DEFAULT '',
  path          TEXT    NOT NULL DEFAULT '',
  kind          TEXT    NOT NULL DEFAULT '',
  size          INTEGER NOT NULL DEFAULT 0,
  username      TEXT    NOT NULL DEFAULT '',
  created_at    TEXT    NOT NULL DEFAULT '',
  snapshot_json TEXT
);
CREATE INDEX IF NOT EXISTS ix_gen_path ON generated_files(path);

CREATE TABLE IF NOT EXISTS meta (
  k TEXT PRIMARY KEY,
  v TEXT NOT NULL DEFAULT ''
);
"""

_init_lock = threading.Lock()
# 记住的是「哪个文件已经建过表」，不是一个 True/False —— 诊断脚本会在同一个进程里
# 把台账指到另一个临时目录，用布尔记就会「以为建过了」，新库一张表都没有。
_initialized_path = None


def init(force: bool = False):
    """建表 + 建索引（幂等，任何入口先调一次就行）"""
    global _initialized_path
    want = str(db_path())
    with _init_lock:
        if _initialized_path == want and not force:
            return
        conn = _open()
        try:
            conn.executescript(_SCHEMA)
            _ensure_unique_indexes(conn)
        finally:
            conn.close()
        _initialized_path = want


def _ensure_unique_indexes(conn):
    """发票号码 / 内容指纹建唯一索引。

    文档要求「发票号码和无号码凭证指纹建立唯一索引」。但历史台账里真的出现过
    「同一号码两行」（早年在两个批次里各入过一次，留痕的），这种情况下唯一索引
    建不上 —— 这时**降级成普通索引**并记一笔，绝不因为索引建不上就让整个程序起不来。
    查重本来就还有应用层那一层，功能不受影响。

    `dupe = 0` 那个条件：整表导入（从 Excel 迁）时如果真撞上重复，那一条会被标成
    dupe=1 保留下来（数据不丢），只是不再参与唯一性约束 —— 见 insert_row(tolerant=True)。
    """
    for name, col in (("ux_inv_no_alive", "inv_no"), ("ux_fp_alive", "fp")):
        try:
            conn.execute(f"CREATE UNIQUE INDEX IF NOT EXISTS {name} "
                         f"ON invoices({col}) WHERE {col} <> '' AND is_deleted = 0 AND dupe = 0")
        except sqlite3.IntegrityError as e:
            conn.execute(f"CREATE INDEX IF NOT EXISTS ix_{name}_fallback ON invoices({col})")
            _set_meta_raw(conn, f"unique_index_{col}_failed", f"{_now()} {e}")


# ======================================================================
# meta
# ======================================================================
def _set_meta_raw(conn, k: str, v: str):
    conn.execute("INSERT INTO meta(k, v) VALUES(?, ?) "
                 "ON CONFLICT(k) DO UPDATE SET v = excluded.v", (k, str(v)))


def meta_get(k: str, default: str = "") -> str:
    init()
    with read() as conn:
        r = conn.execute("SELECT v FROM meta WHERE k = ?", (k,)).fetchone()
    return r["v"] if r else default


def meta_set(k: str, v: str):
    init()
    with write_tx() as conn:
        _set_meta_raw(conn, k, v)


def _file_md5(p) -> str:
    """文件的 md5（读不了就返回空串）—— 用来判断「这份 Excel 还是不是我导出的那份」"""
    try:
        return hashlib.md5(Path(p).read_bytes()).hexdigest()
    except OSError:
        return ""


# ======================================================================
# 台账行：读
# ======================================================================
def _to_dict(row: sqlite3.Row, seq: int) -> dict:
    """数据库行 → 中文列名 dict（和以前 Excel 读出来的一模一样）"""
    d = {f: "" for f in FIELDS}
    d["序号"] = seq
    for cn, dbc in COL2DB.items():
        v = row[dbc]
        d[cn] = "" if v is None else v
    return d


def all_rows() -> list:
    """全部未删除的行，按入库顺序编号 1..N（序号是算出来的，不存库）"""
    init()
    with read() as conn:
        rs = conn.execute("SELECT * FROM invoices WHERE is_deleted = 0 ORDER BY id").fetchall()
    return [_to_dict(r, i) for i, r in enumerate(rs, start=1)]


def all_rows_with_id() -> list:
    """同 all_rows，但额外带 _id（审计/精确定位用；界面不认这个键）"""
    init()
    with read() as conn:
        rs = conn.execute("SELECT * FROM invoices WHERE is_deleted = 0 ORDER BY id").fetchall()
    out = []
    for i, r in enumerate(rs, start=1):
        d = _to_dict(r, i)
        d["_id"] = r["id"]
        out.append(d)
    return out


def count() -> int:
    init()
    with read() as conn:
        return conn.execute("SELECT COUNT(*) c FROM invoices WHERE is_deleted = 0").fetchone()["c"]


# ======================================================================
# 台账行：写
# ======================================================================
def _values_from_row(row: dict) -> dict:
    """中文 dict → 数据库列值"""
    vals = {}
    for cn, dbc in COL2DB.items():
        vals[dbc] = row.get(cn, "")
    for dbc in MONEY_DB:
        vals[dbc] = _num(vals.get(dbc))
    return vals


def _num(v):
    if v in (None, ""):
        return None
    try:
        return round(float(str(v).replace(",", "")), 2)
    except (TypeError, ValueError):
        return None


def insert_row(conn, row: dict, fp: str = "", path_key: str = "",
               tolerant: bool = False) -> int:
    """插一行，返回新 id（conn 由调用方给，必须在事务里）。

    tolerant：整表导入（从 Excel 迁）时用。万一撞上号码/指纹的唯一索引，就退一步把这条
    标成 dupe=1（不再参与唯一性约束）重新插进去。历史台账里确实有「同一张票入过两次」的
    留痕，不能因为一条就让整批导入失败 —— 更不能把这条悄悄丢掉。
    """
    vals = _values_from_row(row)
    vals["fp"] = fp or ""
    vals["path_key"] = path_key or ""
    vals["updated_at"] = _now()
    vals["updated_by"] = actor()
    cols = list(vals.keys())
    sql = (f"INSERT INTO invoices({','.join(cols)}) "
           f"VALUES({','.join('?' for _ in cols)})")
    try:
        cur = conn.execute(sql, [vals[c] for c in cols])
        return cur.lastrowid
    except sqlite3.IntegrityError:
        if not tolerant:
            raise
        vals["dupe"] = 1
        cols2 = list(vals.keys())
        sql2 = (f"INSERT INTO invoices({','.join(cols2)}) "
                f"VALUES({','.join('?' for _ in cols2)})")
        cur = conn.execute(sql2, [vals[c] for c in cols2])
        return cur.lastrowid


def update_row_by_id(conn, rid: int, fields_cn: dict, expect_rev=None) -> int:
    """按 id 改几个字段（字段名是中文列名），返回改了几行。

    expect_rev：乐观锁。传了「我看到的版本号」就要求它和库里一致，不一致这一行就不改、
    返回 0 —— 「两个人同时改同一行时只有一个成功」就是靠它。
    不传就退化成「最后写入生效」；这两种情况下 before/after 都会进审计，改了什么查得回来。
    """
    sets, vals = [], []
    for cn, v in fields_cn.items():
        dbc = COL2DB.get(cn)
        if not dbc:
            continue
        vals.append(_num(v) if dbc in MONEY_DB else ("" if v is None else v))
        sets.append(f"{dbc} = ?")
    if not sets:
        return 0
    vals += [_now(), actor()]
    sets += ["updated_at = ?", "updated_by = ?", "rev = rev + 1"]
    where = "id = ?"
    vals.append(rid)
    if expect_rev is not None:
        where += " AND rev = ?"
        vals.append(int(expect_rev))
    cur = conn.execute(f"UPDATE invoices SET {', '.join(sets)} WHERE {where}", vals)
    return cur.rowcount


def seq_to_id(conn, seqs) -> list:
    """界面给的「序号」→ 行 id。

    序号是「按 id 排序后的第几行」，所以要现算：这里和 all_rows 用的是同一个排序，
    界面看到的第 3 行一定对应这里算出来的第 3 个 id。
    """
    want = {int(s) for s in seqs if str(s).strip().isdigit()}
    if not want:
        return []
    ids = [r["id"] for r in conn.execute(
        "SELECT id FROM invoices WHERE is_deleted = 0 ORDER BY id")]
    return [ids[s - 1] for s in sorted(want) if 0 < s <= len(ids)]


def find_existing(conn, no: str = "", fp: str = "", src: str = "") -> list:
    """查重：先按发票号码，再按内容指纹，最后兜一层「同一个文件路径」。

    返回命中的数据库行列表（同一号码历史上可能有两行）。查重只看**未删除**的行。
    """
    def q(where, val):
        if not val:
            return []
        return conn.execute(
            f"SELECT * FROM invoices WHERE is_deleted = 0 AND {where} = ? ORDER BY id",
            (val,)).fetchall()

    hits = q("inv_no", (no or "").strip())
    if not hits and not (no or "").strip():
        hits = q("fp", fp)
    if not hits:
        hits = q("path_key", norm_path(src))
    return hits


def replace_all(rows: list):
    """整表替换（给「从 Excel 导入」和测试准备数据用）。

    故意保持这个「整表写回」的口子 —— 老的调用方（和一批诊断脚本）就是这么用的。
    真正的并发安全来自 add_records / update_rows 那些走增量 SQL 的路径。
    """
    init()
    now = _now()
    with write_tx() as conn:
        conn.execute("DELETE FROM invoices")
        for row in rows or []:
            row = dict(row)
            row.setdefault("入库时间", now)
            fp = _fp_of(row)
            insert_row(conn, row, fp=fp, path_key=norm_path(row.get("文件路径")),
                       tolerant=True)


def norm_path(p) -> str:
    """路径归一化（大小写、斜杠方向都可能不一样）"""
    return str(p or "").replace("\\", "/").strip().lower()


def _fp_of(row: dict) -> str:
    """台账行的内容指纹 —— 只对**没有发票号码**的凭证有意义。

    ⚠️ 有号码的凭证靠号码查重，不能再占指纹索引：否则「两张号码不同、但类型/日期/金额/
    明细恰好一样」的票（同一天同一个价位的两张打车票）会撞上唯一索引，整批导入直接失败。
    口径和 invparse.fingerprint 一致，跟入库时算的是同一个值。
    """
    if str(row.get("发票号码") or "").strip():
        return ""
    if not any(str(row.get(k) or "") for k in ("开票日期", "价税合计", "行程/明细", "文件路径")):
        return ""
    try:
        import invparse
    except ImportError:
        return ""
    return invparse.fingerprint({
        "ctype": row.get("凭证类型", ""),
        "date": row.get("开票日期", ""),
        "total": row.get("价税合计"),
        "detail": row.get("行程/明细", ""),
        "src": row.get("文件路径", ""),
    })


def soft_delete(conn, ids, reason: str = "") -> list:
    """软删除：只是标记一下，行还在库里（审计要查得回来），界面和查重都看不见了"""
    out = []
    for rid in ids:
        r = conn.execute("SELECT * FROM invoices WHERE id = ?", (rid,)).fetchone()
        if r is None or r["is_deleted"]:
            continue
        conn.execute("UPDATE invoices SET is_deleted = 1, deleted_at = ?, deleted_by = ? "
                     "WHERE id = ?", (_now(), actor(), rid))
        out.append(dict(r))
    return out


def clear_all(conn, reason: str = "") -> int:
    """清空（软删除所有行）—— 老代码里 clear_ledger 的语义：清完表就空了"""
    rows = conn.execute("SELECT COUNT(*) c FROM invoices WHERE is_deleted = 0").fetchone()["c"]
    conn.execute("UPDATE invoices SET is_deleted = 1, deleted_at = ?, deleted_by = ? "
                 "WHERE is_deleted = 0", (_now(), actor()))
    return rows


# ======================================================================
# 审计日志
# ======================================================================
def audit(action: str, object_type: str = "", object_id: str = "",
          before=None, after=None, ok: bool = True, error: str = "",
          username: str = None, ip: str = None):
    """写一条审计。

    刻意**不**参与调用方的事务：调用方的事务回滚了，「这次操作失败」这件事本身也得留下。
    所以自己开一个连接写完就关。
    """
    init()
    try:
        before_j = json.dumps(before, ensure_ascii=False, default=str) if before is not None else None
        after_j = json.dumps(after, ensure_ascii=False, default=str) if after is not None else None
        with write_tx() as conn:
            conn.execute(
                "INSERT INTO audit_logs(ts, username, ip, action, object_type, object_id,"
                " before_json, after_json, ok, error) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (_now(), username if username is not None else actor(),
                 ip if ip is not None else client_ip(), action, object_type or "",
                 str(object_id or ""), before_j, after_j, 1 if ok else 0, error or ""))
    except Exception:
        # 审计写不进去绝不能拖垮主流程（磁盘满、库损坏都不该让用户没法入库）
        pass


def list_audit(limit: int = 200, action: str = "", username: str = "",
               object_id: str = "", offset: int = 0) -> list:
    init()
    where, vals = [], []
    if action:
        where.append("action = ?")
        vals.append(action)
    if username:
        where.append("username = ?")
        vals.append(username)
    if object_id:
        where.append("object_id = ?")
        vals.append(str(object_id))
    sql = "SELECT * FROM audit_logs"
    if where:
        sql += " WHERE " + " AND ".join(where)
    sql += " ORDER BY id DESC LIMIT ? OFFSET ?"
    vals += [int(limit), int(offset)]
    with read() as conn:
        return [dict(r) for r in conn.execute(sql, vals).fetchall()]


def count_audit() -> int:
    init()
    with read() as conn:
        return conn.execute("SELECT COUNT(*) c FROM audit_logs").fetchone()["c"]


# ======================================================================
# 后台任务
# ======================================================================
JOB_QUEUED, JOB_RUNNING, JOB_DONE, JOB_FAILED, JOB_CANCELLED = \
    "排队", "运行中", "成功", "失败", "已取消"


def job_create(kind: str, label: str = "", total: int = 0) -> str:
    init()
    jid = secrets.token_hex(8)
    with write_tx() as conn:
        conn.execute("INSERT INTO jobs(id, kind, state, total, done, label, username, created_at)"
                     " VALUES(?,?,?,?,0,?,?,?)",
                     (jid, kind, JOB_QUEUED, int(total), label or "", actor(), _now()))
    return jid


def job_update(jid: str, **f):
    init()
    allow = {"state", "total", "done", "label", "error", "result_json",
             "started_at", "ended_at"}
    sets, vals = [], []
    for k, v in f.items():
        if k not in allow:
            continue
        if k == "result_json" and not isinstance(v, (str, type(None))):
            v = json.dumps(v, ensure_ascii=False, default=str)
        sets.append(f"{k} = ?")
        vals.append(v)
    if not sets:
        return
    vals.append(jid)
    with write_tx() as conn:
        conn.execute(f"UPDATE jobs SET {', '.join(sets)} WHERE id = ?", vals)


def job_get(jid: str) -> dict:
    init()
    with read() as conn:
        r = conn.execute("SELECT * FROM jobs WHERE id = ?", (jid,)).fetchone()
    return dict(r) if r else {}


def job_list(limit: int = 50) -> list:
    init()
    with read() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM jobs ORDER BY created_at DESC, rowid DESC LIMIT ?",
            (int(limit),)).fetchall()]


def job_orphans_fix() -> int:
    """进程重启后，把「上一世」还挂着的任务标成失败 —— 否则界面会一直转圈等一个不存在的任务"""
    init()
    with write_tx() as conn:
        cur = conn.execute(
            "UPDATE jobs SET state = ?, ended_at = ?, error = ? "
            "WHERE state IN (?, ?)",
            (JOB_FAILED, _now(), "服务重启，任务中断", JOB_QUEUED, JOB_RUNNING))
        return cur.rowcount


# ======================================================================
# 生成文件
# ======================================================================
def gen_file_add(path: str, kind: str = "", size: int = 0, snapshot=None, name: str = ""):
    init()
    p = Path(path)
    with write_tx() as conn:
        conn.execute(
            "INSERT INTO generated_files(name, path, kind, size, username, created_at, snapshot_json)"
            " VALUES(?,?,?,?,?,?,?)",
            (name or p.name, str(p), kind or "", int(size or 0), actor(), _now(),
             json.dumps(snapshot, ensure_ascii=False, default=str) if snapshot is not None else None))


def gen_file_list(limit: int = 200) -> list:
    init()
    with read() as conn:
        return [dict(r) for r in conn.execute(
            "SELECT * FROM generated_files ORDER BY id DESC LIMIT ?", (int(limit),)).fetchall()]


# ======================================================================
# 用户与口令
# ======================================================================
ROLES = {"admin": "管理员", "auditor": "财务审核员", "biz": "业务人员", "viewer": "只读用户"}
ROLE_ALLOW = {
    # 每种角色能调的后端方法（app_web 里逐个接口校验用的白名单）
    "admin":   None,                                  # None = 全部允许
    "auditor": {"get_ledger", "mark_rows", "update_rows", "refresh_rows",
                "traveler_review", "save_traveler_review", "get_phone_map", "set_phone_name",
                "make_report", "report_persons", "merge_invoices", "export_summary",
                "list_outputs", "delete_output", "open_folder", "open_file", "reveal",
                "upload_info", "get_paths", "get_config", "whoami", "ledger_db_info", "poll", "list_jobs",
                "audit_list", "gen_files"},
    "biz":     {"get_ledger", "mark_rows", "update_rows", "traveler_review",
                "save_traveler_review", "get_phone_map", "set_phone_name",
                "make_report", "report_persons", "merge_invoices", "export_summary",
                "list_outputs", "open_folder", "open_file", "reveal", "upload_info",
                "check_source", "start_import", "get_paths", "get_config", "whoami", "ledger_db_info", "poll",
                "list_jobs", "gen_files"},
    "viewer":  {"get_ledger", "get_paths", "get_config", "whoami", "ledger_db_info", "poll", "list_outputs",
                "gen_files"},
}


# ---------------------------------------------------------------- 口令哈希
# 优先 Argon2id（文档点名的）；环境里没装 argon2-cffi 就退回标准库的 scrypt。
# 两者都是「慢哈希」，都能用；存下来的字符串自带前缀，校验时按前缀分派。
def hash_password(pw: str) -> str:
    try:
        from argon2 import PasswordHasher
        return PasswordHasher().hash(pw)
    except ImportError:
        pass
    salt = secrets.token_bytes(16)
    n, r, p = 2 ** 14, 8, 1
    dk = hashlib.scrypt(pw.encode("utf-8"), salt=salt, n=n, r=r, p=p, dklen=32)
    return f"scrypt${n}${r}${p}${salt.hex()}${dk.hex()}"


def verify_password(pw: str, stored: str) -> bool:
    if not stored:
        return False
    if stored.startswith("$argon2"):
        try:
            from argon2 import PasswordHasher
            from argon2.exceptions import VerifyMismatchError
            try:
                PasswordHasher().verify(stored, pw)
                return True
            except VerifyMismatchError:
                return False
            except Exception:
                return False
        except ImportError:
            return False
    if stored.startswith("scrypt$"):
        try:
            _, n, r, p, salt_hex, dk_hex = stored.split("$")
            dk = hashlib.scrypt(pw.encode("utf-8"), salt=bytes.fromhex(salt_hex),
                                n=int(n), r=int(r), p=int(p), dklen=len(dk_hex) // 2)
            return hmac.compare_digest(dk.hex(), dk_hex)
        except Exception:
            return False
    # 明文字令（老的 web_password 配置项）—— 只读校验，不再产生新的。
    # ⚠️ compare_digest 不能拿含中文的 str 直接比，会 TypeError → 统一转成字节再比
    return hmac.compare_digest(pw.encode("utf-8"), stored.encode("utf-8"))


def user_create(username: str, password: str, role: str = "biz",
                display_name: str = "", dept: str = "", scope: str = "") -> int:
    init()
    if role not in ROLES:
        raise ValueError(f"角色必须是 {list(ROLES)} 之一")
    with write_tx() as conn:
        cur = conn.execute(
            "INSERT INTO users(username, display_name, pw_hash, role, dept, scope, created_at)"
            " VALUES(?,?,?,?,?,?,?)",
            (username.strip(), display_name or "", hash_password(password), role,
             dept or "", scope or "", _now()))
        return cur.lastrowid


def user_get(username: str) -> dict:
    init()
    with read() as conn:
        r = conn.execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()
    return dict(r) if r else {}


def user_list() -> list:
    init()
    with read() as conn:
        return [dict(r) for r in conn.execute("SELECT * FROM users ORDER BY id").fetchall()]


def user_count() -> int:
    init()
    with read() as conn:
        return conn.execute("SELECT COUNT(*) c FROM users").fetchone()["c"]


def user_set_disabled(username: str, disabled: bool):
    init()
    with write_tx() as conn:
        conn.execute("UPDATE users SET disabled = ? WHERE username = ?",
                     (1 if disabled else 0, username))


def user_set_password(username: str, password: str):
    init()
    with write_tx() as conn:
        conn.execute("UPDATE users SET pw_hash = ? WHERE username = ?",
                     (hash_password(password), username))


def user_set_role(username: str, role: str):
    init()
    if role not in ROLES:
        raise ValueError(f"角色必须是 {list(ROLES)} 之一")
    with write_tx() as conn:
        conn.execute("UPDATE users SET role = ? WHERE username = ?", (role, username))


def user_touch_login(username: str):
    init()
    with write_tx() as conn:
        conn.execute("UPDATE users SET last_login = ? WHERE username = ?", (_now(), username))


def can(username: str, role: str, method: str) -> bool:
    """这个方法这个角色能不能调（None = 管理员，全放开）"""
    allow = ROLE_ALLOW.get(role, set())
    if allow is None:
        return True
    return method in allow


# ======================================================================
# 从老 Excel 迁移
# ======================================================================
def read_excel_rows(excel_path) -> list:
    """把老台账 Excel 读成中文 dict 列表（表头认不出来就当空表）"""
    import openpyxl
    p = Path(excel_path)
    if not p.exists():
        return []
    wb = openpyxl.load_workbook(str(p), data_only=True)
    ws = wb["发票台账"] if "发票台账" in wb.sheetnames else wb[wb.sheetnames[0]]
    header = [(_c.value or "").strip() if isinstance(_c.value, str) else _c.value
              for _c in ws[1]]
    if not header or header[0] != "序号":
        return []
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r is None or all(v in (None, "") for v in r):
            continue
        row = {f: "" for f in FIELDS}
        for i, name in enumerate(header):
            if i < len(r) and name in row:
                row[name] = r[i] if r[i] is not None else ""
        rows.append(row)
    return rows


def migrate_from_excel(excel_path, force: bool = False) -> dict:
    """老 Excel → SQLite。

    规矩（文档第 1.5 条）：迁移前先把原 Excel 备份一份，然后**只往里加、不删原文件**。
    已经迁过的（meta 里有标记 且 库里已有行）默认不再迁第二次，除非 force。
    """
    init()
    p = Path(excel_path)
    result = {"ok": False, "migrated": 0, "backup": "", "skipped": "", "excel": str(p)}
    if not p.exists():
        result["skipped"] = "没有找到 Excel 台账文件"
        return result
    if not force and count() > 0:
        # 库里已经有数据了，默认什么都不做。
        # ⚠️ 自动导入 = 整表替换，这种事绝不能悄悄发生：用户已经在界面上改了半天，
        #    启动时被一份旧的 Excel 覆盖掉就完了。想覆盖必须显式点「从 Excel 导入」。
        if _file_md5(p) == meta_get("excel_export_hash"):
            result.update({"ok": True, "skipped": "Excel 就是本程序导出的快照，无需迁移"})
        elif meta_get("migrated_from_excel"):
            result.update({"ok": True, "skipped": "已经迁移过，跳过（要重来请用 force）"})
        else:
            result.update({"ok": True,
                           "skipped": "库里已有数据；这份 Excel 和程序导出的快照对不上，"
                                      "如需并回请用「从 Excel 导入」"})
        return result
    rows = read_excel_rows(p)
    if not rows:
        result["skipped"] = "Excel 里没有可识别的数据行"
        result["ok"] = True
        return result

    # ① 先备份原文件（同一秒撞名就加序号，别把上一份覆盖了）
    bak_dir = Path(app_paths.LEDGER_BAK)
    bak_dir.mkdir(parents=True, exist_ok=True)
    stem = f"发票台账_迁移前_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    dst = bak_dir / f"{stem}.xlsx"
    n = 1
    while dst.exists() and n < 100:
        dst = bak_dir / f"{stem}_{n}.xlsx"
        n += 1
    import shutil
    shutil.copy2(p, dst)
    result["backup"] = str(dst)

    # ② 事务里整批写入
    now = _now()
    with write_tx() as conn:
        conn.execute("DELETE FROM invoices")
        for row in rows:
            row.setdefault("入库时间", now)
            insert_row(conn, row, fp=_fp_of(row), path_key=norm_path(row.get("文件路径")),
                       tolerant=True)
        _set_meta_raw(conn, "migrated_from_excel", f"{now} <- {p.name}")
        _set_meta_raw(conn, "migrated_rows", str(len(rows)))
    result.update({"ok": True, "migrated": len(rows)})
    audit("迁移台账", "ledger", p.name if p else "",
          before={"excel": str(p)}, after={"rows": len(rows), "backup": str(dst)})
    return result


# ======================================================================
# 自检
# ======================================================================
def stats() -> dict:
    init()
    with read() as conn:
        inv = conn.execute("SELECT COUNT(*) c FROM invoices WHERE is_deleted = 0").fetchone()["c"]
        dele = conn.execute("SELECT COUNT(*) c FROM invoices WHERE is_deleted = 1").fetchone()["c"]
        au = conn.execute("SELECT COUNT(*) c FROM audit_logs").fetchone()["c"]
        jo = conn.execute("SELECT COUNT(*) c FROM jobs").fetchone()["c"]
        us = conn.execute("SELECT COUNT(*) c FROM users").fetchone()["c"]
    p = db_path()
    try:
        size = p.stat().st_size
    except OSError:
        size = 0
    return {"db": str(p), "size": size, "invoices": inv, "deleted": dele,
            "audit": au, "jobs": jo, "users": us}


if __name__ == "__main__":
    import sys
    if len(sys.argv) > 1 and sys.argv[1] == "migrate":
        from app_paths import LEDGER_FILE
        print(json.dumps(migrate_from_excel(sys.argv[2] if len(sys.argv) > 2 else LEDGER_FILE),
                         ensure_ascii=False, indent=2))
    else:
        init()
        print(json.dumps(stats(), ensure_ascii=False, indent=2))
